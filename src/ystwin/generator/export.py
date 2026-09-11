"""Write a synthetic plate in the shape the plate reader writes.

Round-tripping through the real export format is the generator's acceptance test: a plate that
needs a special reader is not standing in for a wet-lab file.
"""

from __future__ import annotations

import pathlib

import pandas as pd
from openpyxl import Workbook

from ..qpcr import STRESSOR_FOR_CONSTRUCT
from .plate import SyntheticPlate

__all__ = ["derived_dose_response", "write_synergy_export"]


def _clock(hours: float) -> str:
    total = int(round(hours * 3600))
    return f"{total // 3600:02d}:{(total % 3600) // 60:02d}:{total % 60:02d}"


def write_synergy_export(
    plate: SyntheticPlate,
    path: str | pathlib.Path,
    include_dose_sheets: bool = True,
) -> pathlib.Path:
    """Write a synthetic plate as a Synergy-style workbook.

    Args:
        plate: Generated plate.
        path: Destination ``.xlsx``.
        include_dose_sheets: Also write the per-construct dose-response sheets the
            team produces by hand, so the layout recovery has something to match.
    """
    path = pathlib.Path(path)
    workbook = Workbook()
    workbook.remove(workbook.active)

    for sheet_name, frame in (("OD600", plate.od), ("mCitrine", plate.rfu)):
        sheet = workbook.create_sheet(sheet_name)
        sheet.append([])
        sheet.append([])
        sheet.append(["Time", *frame.columns])
        for time_h, row in frame.iterrows():
            sheet.append([_clock(float(time_h)), *[float(v) for v in row]])

    if include_dose_sheets:
        derived = derived_dose_response(plate)
        for construct, group in derived.groupby("construct"):
            sheet = workbook.create_sheet(str(construct))
            sheet.append([])
            doses = sorted(group.dose_mM.unique())
            sheet.append([None, "Time", *[f"{d:g} mM" for d in doses]])
            wide = group.pivot_table(index="time_h", columns="dose_mM", values="signal")
            for time_h, row in wide.iterrows():
                sheet.append([None, float(time_h) * 60.0, *[float(row[d]) for d in doses]])

    workbook.save(path)
    workbook.close()
    return path


def derived_dose_response(plate: SyntheticPlate) -> pd.DataFrame:
    """The per-construct blanked signal, as the team computes it by hand.

    ``(RFU - background) / (OD - blank)``, averaged over each condition's replicate
    wells -- the same quantity their own sheets hold, so layout recovery can be
    exercised on synthetic plates exactly as on real ones.
    """
    blanks = set(plate.blank_wells)
    background = plate.conditions.optics.background
    od_blank = plate.conditions.od_blank

    rows = []
    truth = plate.truth[~plate.truth.well.isin(blanks)]
    for (construct, dose), group in truth.groupby(["construct", "dose_mM"]):
        wells = [w for w in group.well if w in plate.od.columns]
        if not wells:
            continue
        signal = (
            (plate.rfu[wells] - background) / (plate.od[wells] - od_blank)
        ).mean(axis=1)
        for time_h, value in signal.items():
            rows.append({
                "construct": construct,
                "stressor": STRESSOR_FOR_CONSTRUCT.get(construct, "DTT"),
                "dose_mM": float(dose),
                "time_h": float(time_h),
                "signal": float(value),
            })
    return pd.DataFrame(rows)
