from __future__ import annotations

from collections import Counter
from dataclasses import asdict, dataclass, field, replace
from enum import Enum
import hashlib
from io import BytesIO
import math
from pathlib import Path
import posixpath
from types import MappingProxyType
from typing import Iterator, Literal, Mapping, Sequence, TypeAlias
import xml.etree.ElementTree as ET
from zipfile import BadZipFile, ZipFile

import openpyxl
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries
import pandas as pd


__all__ = [
    "CosmicCell", "CosmicDataError", "CosmicDataset", "CosmicQuantity", "CosmicSheet",
    "CosmicSource", "EvidenceRole", "QuantityRole", "VesselMapping",
    "PUBLISHER_WORKBOOK_BYTES", "PUBLISHER_WORKBOOK_SHA256", "SYNTHETIC_WORKBOOK_TITLE",
    "WORKBOOK_FILENAMES", "load_cosmic_workbook", "parse_synthetic_cosmic_workbook",
    "summarize_cosmic_data",
]

PUBLISHER_WORKBOOK_SHA256 = "a68dc8bf86b87e6b9d4f732cd3413948bdb567bf3f85fdabe937a202bb7a7daa"
PUBLISHER_WORKBOOK_BYTES = 4156209
WORKBOOK_FILENAMES = (
    "1-s2.0-S1096717624000284-mmc2.xlsx", "COSMIC-dFBA_supplementary_data.xlsx",
)
SYNTHETIC_WORKBOOK_TITLE = "SYNTHETIC COSMIC parser fixture - not scientific evidence"
CellValue: TypeAlias = str | int | float | None

_SHEET_NAMES = ("Index", "ST1", "ST2", "ST3", "ST4", "ST5")
_SHAPES = {"Index": (6, 2), "ST1": (14, 4), "ST2": (132, 28), "ST3": (27, 22), "ST4": (13, 41), "ST5": (4321, 167)}
_NONEMPTY_CELLS = {"Index": 10, "ST1": 46, "ST2": 3695, "ST3": 573, "ST4": 417, "ST5": 721607}
_VESSELS = ("R0001", "R0002", "R0003", "R0004", "R0005", "R0006", "R0008", "R0010", "R0011", "R0012")
_PHASES = ("Growth Phase", "Production Phase")
_PROCESS_LABELS = (
    "Vessel", "Time", "Production phase fraction", "Cell Density", "Cell Volume",
    "Glucose", "Lactate", "NH4", "Titer", "Glutamine", "Glutamate", "L-Asparagine",
    "L-Aspartic acid", "L-Serine", "Glycine", "L-Alanine ", "L-Proline", "L-Threonine",
    "L-Histidine", "L-Lysine", "L-Valine ", "L-Methionine", "L-Arginine", "L-Tyrosine",
    "L-Isoleucine", "L-Leucine", "L-Phenylalanine", "L-Tryptophan",
)
_INDEX_DESCRIPTIONS = (
    "Box Behnken design with 3 factors",
    "Normalized measured process data",
    "Computed uptake and secretion rates",
    "Cellular objectives in the growth and production phases",
    "Intracellular fluxes (mmol/gDW-d) computed using COSMIC-dFBA",
)
_DESIGN_NOTE = (
    "*0 represents a central value, +1 and -1 represent increased and reduced values "
    "for the independent variable, respectively"
)
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
_NS = {"s": _MAIN_NS}


class CosmicDataError(ValueError):
    pass


class QuantityRole(str, Enum):
    CODED_DESIGN_LEVEL = "coded_design_level"
    NORMALIZED_PROCESS_QUANTITY = "normalized_process_quantity"
    PRODUCTION_PHASE_FRACTION = "production_phase_fraction"
    COMPUTED_RATE = "computed_rate"
    OBJECTIVE_EFFICIENCY = "objective_efficiency"
    INTRACELLULAR_FLUX = "intracellular_flux"


class EvidenceRole(str, Enum):
    SOURCE_DESIGN = "source_design"
    NORMALIZED_MEASUREMENT = "normalized_source_measurement"
    STATE_ANNOTATION = "source_state_annotation_not_independent_measurement"
    COMPUTED_RATE = "computed_source_rate"
    FITTED_OBJECTIVE = "fitted_source_objective"
    COMPUTED_FLUX = "computed_source_flux"


@dataclass(frozen=True, slots=True)
class CosmicSource:
    kind: Literal["publisher_workbook", "synthetic_fixture"]
    sha256: str
    byte_count: int
    doi: str | None
    host: str


@dataclass(frozen=True, slots=True)
class CosmicCell:
    source: CosmicSource
    sheet: str
    coordinate: str
    value: CellValue

    @property
    def identity(self) -> str:
        return f"{self.source.kind}:{self.source.sha256}:{self.sheet}!{self.coordinate}"


@dataclass(frozen=True, slots=True)
class CosmicQuantity:
    cell: CosmicCell
    label: CosmicCell
    quantity_role: QuantityRole
    evidence_role: EvidenceRole
    role_evidence: CosmicCell
    source_unit: str | None
    unit_cell: CosmicCell | None
    normalization: str
    normalization_cell: CosmicCell | None = None
    vessel: CosmicCell | None = None
    phase: CosmicCell | None = None
    time: CosmicCell | None = None
    source_time_unit: str | None = None
    time_unit_cell: CosmicCell | None = None
    priority: CosmicCell | None = None
    state_fraction: CosmicCell | None = None

    @property
    def value(self) -> int | float | None:
        return self.cell.value

    @property
    def is_missing(self) -> bool:
        return self.cell.value is None


@dataclass(frozen=True, slots=True)
class CosmicSheet:
    source: CosmicSource
    name: str
    values: tuple[tuple[CellValue, ...], ...] = field(repr=False)
    merged_ranges: tuple[str, ...]

    @property
    def shape(self) -> tuple[int, int]:
        return len(self.values), len(self.values[0])

    @property
    def blank_rows(self) -> tuple[int, ...]:
        return tuple(r for r, row in enumerate(self.values, 1) if all(v is None for v in row))

    @property
    def nonempty_row_count(self) -> int:
        return self.shape[0] - len(self.blank_rows)

    @property
    def nonempty_cell_count(self) -> int:
        return sum(value is not None for row in self.values for value in row)

    def cell(self, coordinate: str) -> CosmicCell:
        row, column = _coordinate(coordinate, self.name)
        if row > self.shape[0] or column > self.shape[1]:
            raise CosmicDataError(f"Source cell outside {self.name} shape: {coordinate}")
        return CosmicCell(self.source, self.name, coordinate, self.values[row - 1][column - 1])

    def at(self, row: int, column: int) -> CosmicCell:
        if type(row) is not int or type(column) is not int or row < 1 or column < 1:
            raise CosmicDataError("Source cell indices must be positive integers")
        return self.cell(f"{get_column_letter(column)}{row}")

    def to_frame(self) -> pd.DataFrame:
        return pd.DataFrame(
            self.values, index=range(1, self.shape[0] + 1),
            columns=[get_column_letter(c) for c in range(1, self.shape[1] + 1)], dtype=object,
        )


@dataclass(frozen=True, slots=True)
class VesselMapping:
    matched: tuple[str, ...]
    workbook_only: tuple[str, ...]
    target_only: tuple[str, ...]
    method: str = "exact_identifier_only_no_ordinal_mapping"


@dataclass(frozen=True, slots=True)
class CosmicDataset:
    source: CosmicSource
    sheets: Mapping[str, CosmicSheet] = field(repr=False)
    sha256_after: str

    def cell(self, sheet: str, coordinate: str) -> CosmicCell:
        if sheet not in self.sheets:
            raise CosmicDataError(f"Unknown source sheet: {sheet}")
        return self.sheets[sheet].cell(coordinate)

    def vessel_ids(self, sheet: str) -> tuple[str, ...]:
        if sheet not in self.sheets:
            raise CosmicDataError(f"Unknown source sheet: {sheet}")
        values = self.sheets[sheet].values
        if sheet == "ST1":
            return tuple(row[0] for row in values[2:12])
        if sheet == "ST2":
            return tuple(dict.fromkeys(row[0] for row in values[2:]))
        if sheet == "ST3":
            return tuple(dict.fromkeys(values[1][2:]))
        if sheet == "ST4":
            return values[0][1::4]
        return ()

    def match_vessels(self, target_ids: Sequence[str]) -> VesselMapping:
        if isinstance(target_ids, (str, bytes)):
            raise CosmicDataError("Target vessel identities must be a sequence, not a string")
        targets = tuple(_identifier(value, "target vessel") for value in target_ids)
        _unique(targets, "target vessel")
        return _match_vessels(self.vessel_ids("ST1"), targets)

    def quantity(self, sheet: str, coordinate: str) -> CosmicQuantity:
        cell = self.cell(sheet, coordinate)
        row, column = _coordinate(coordinate, sheet)
        table = self.sheets[sheet]
        if sheet == "ST1" and 3 <= row <= 12 and 2 <= column <= 4:
            return CosmicQuantity(
                cell, table.at(2, column), QuantityRole.CODED_DESIGN_LEVEL,
                EvidenceRole.SOURCE_DESIGN, self.cell("Index", "B2"), None, None,
                "coded_relative_levels_no_physical_mapping", self.cell("ST1", "A14"),
                vessel=table.at(row, 1),
            )
        if sheet == "ST2" and row >= 3 and column >= 3:
            annotation = column == 3
            return CosmicQuantity(
                cell, table.at(1, column),
                QuantityRole.PRODUCTION_PHASE_FRACTION if annotation else QuantityRole.NORMALIZED_PROCESS_QUANTITY,
                EvidenceRole.STATE_ANNOTATION if annotation else EvidenceRole.NORMALIZED_MEASUREMENT,
                self.cell("Index", "B3"), table.at(2, column).value, table.at(2, column),
                "source_state_annotation_method_not_supplied" if annotation else "source_normalized_scale_not_supplied",
                self.cell("Index", "B3"), vessel=table.at(row, 1), time=table.at(row, 2),
                source_time_unit=table.at(2, 2).value, time_unit_cell=table.at(2, 2),
                state_fraction=table.at(row, 3),
            )
        if sheet == "ST3" and row >= 3 and column >= 3:
            return CosmicQuantity(
                cell, table.at(row, 1), QuantityRole.COMPUTED_RATE, EvidenceRole.COMPUTED_RATE,
                self.cell("Index", "B4"), table.at(row, 2).value, table.at(row, 2),
                "source_rate_units_no_conversion", vessel=table.at(2, column),
                phase=table.at(1, 3 if column <= 12 else 13),
            )
        if sheet == "ST4" and row >= 4 and column >= 3 and column % 2 == 1:
            group_column = 2 + 4 * ((column - 2) // 4)
            return CosmicQuantity(
                cell, table.at(row, column - 1), QuantityRole.OBJECTIVE_EFFICIENCY,
                EvidenceRole.FITTED_OBJECTIVE, self.cell("Index", "B5"), None, None,
                "source_efficiency_not_clipped_unit_not_stated", vessel=table.at(1, group_column),
                phase=table.at(2, column - 1), priority=table.at(row, 1),
            )
        if sheet == "ST5" and row >= 2 and column >= 2:
            return CosmicQuantity(
                cell, table.at(row, 1), QuantityRole.INTRACELLULAR_FLUX,
                EvidenceRole.COMPUTED_FLUX, self.cell("Index", "B6"), "mmol/gDW-d",
                self.cell("Index", "B6"), "source_flux_units_no_conversion",
                time=table.at(1, column), source_time_unit="days", time_unit_cell=table.at(1, 1),
            )
        raise CosmicDataError(f"Not a quantity cell: {sheet}!{coordinate}")

    def iter_quantities(self, sheet: str) -> Iterator[CosmicQuantity]:
        if sheet not in _SHEET_NAMES[1:]:
            raise CosmicDataError(f"No quantity cells for sheet: {sheet}")
        rows, columns = self.sheets[sheet].shape
        first_row, first_column, step = {
            "ST1": (3, 2, 1), "ST2": (3, 3, 1), "ST3": (3, 3, 1),
            "ST4": (4, 3, 2), "ST5": (2, 2, 1),
        }[sheet]
        last_row = 12 if sheet == "ST1" else rows
        for row in range(first_row, last_row + 1):
            for column in range(first_column, columns + 1, step):
                yield self.quantity(sheet, f"{get_column_letter(column)}{row}")


def _coordinate(value: str, sheet: str) -> tuple[int, int]:
    try:
        row, column = coordinate_to_tuple(value)
        if row < 1 or value != f"{get_column_letter(column)}{row}":
            raise ValueError
    except (AttributeError, KeyError, TypeError, ValueError) as exc:
        raise CosmicDataError(f"Invalid canonical source cell: {sheet}!{value}") from exc
    return row, column


def _identifier(value: CellValue, context: str) -> str:
    if not isinstance(value, str) or not value.strip() or value != value.strip():
        raise CosmicDataError(f"{context}: expected an explicit, unpadded source identifier")
    return value


def _unique(values: Sequence, context: str) -> None:
    if len(values) != len(set(values)):
        raise CosmicDataError(f"Duplicate {context} identities")


def _match_vessels(workbook_ids: tuple[str, ...], target_ids: tuple[str, ...]) -> VesselMapping:
    return VesselMapping(
        tuple(v for v in workbook_ids if v in target_ids),
        tuple(v for v in workbook_ids if v not in target_ids),
        tuple(v for v in target_ids if v not in workbook_ids),
    )


def _expected_merges(sheet: str) -> set[str]:
    if sheet == "ST1":
        return {"B1:D1"}
    if sheet == "ST3":
        return {"C1:L1", "M1:V1"}
    if sheet == "ST4":
        return {
            f"{get_column_letter(c)}1:{get_column_letter(c + 3)}1" for c in range(2, 42, 4)
        } | {
            f"{get_column_letter(c)}2:{get_column_letter(c + 1)}2" for c in range(2, 42, 2)
        }
    return set()


def _sheet_members(archive: ZipFile, synthetic: bool) -> dict[str, str]:
    names = archive.namelist()
    if len(names) != len(set(names)):
        raise CosmicDataError("Duplicate ZIP member names")
    if sum(info.file_size for info in archive.infolist()) > 128 * 1024 * 1024:
        raise CosmicDataError("Workbook ZIP exceeds the reviewed uncompressed size limit")
    if synthetic:
        properties = ET.fromstring(archive.read("docProps/core.xml"))
        titles = properties.findall("{http://purl.org/dc/elements/1.1/}title")
        if len(titles) != 1 or titles[0].text != SYNTHETIC_WORKBOOK_TITLE:
            raise CosmicDataError("Synthetic parsing requires the explicit SYNTHETIC workbook title")
    workbook = ET.fromstring(archive.read("xl/workbook.xml"))
    sheets = workbook.findall("s:sheets/s:sheet", _NS)
    if tuple(sheet.get("name") for sheet in sheets) != _SHEET_NAMES:
        raise CosmicDataError("Unexpected workbook sheet names/order; all six reviewed sheets are required")
    relations = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    rows = relations.findall(f"{{{_PACKAGE_REL_NS}}}Relationship")
    _unique(tuple(row.get("Id") for row in rows), "workbook relationship")
    by_id = {row.get("Id"): row for row in rows}
    members = {}
    for sheet in sheets:
        name = sheet.get("name")
        if sheet.get("state", "visible") != "visible":
            raise CosmicDataError(f"Hidden source sheet is not in the reviewed schema: {name}")
        relationship = by_id.get(sheet.get(f"{{{_REL_NS}}}id"))
        if relationship is None or relationship.get("Type") != f"{_REL_NS}/worksheet":
            raise CosmicDataError(f"Invalid worksheet relationship for {name}")
        if relationship.get("TargetMode", "Internal") != "Internal":
            raise CosmicDataError(f"External worksheet relationship for {name}")
        target = relationship.get("Target", "")
        member = posixpath.normpath(target.lstrip("/") if target.startswith("/") else posixpath.join("xl", target))
        if not member.startswith("xl/worksheets/") or member not in names:
            raise CosmicDataError(f"Missing/nonlocal worksheet member for {name}")
        members[name] = member
    _unique(tuple(members.values()), "worksheet relationship target")
    return members


def _xml_layout(archive: ZipFile, member: str, sheet: str) -> tuple[tuple[int, int], tuple[str, ...]]:
    dimension = None
    current_row = None
    last_row = last_column = max_row = max_column = 0
    merges = []
    with archive.open(member) as stream:
        for event, node in ET.iterparse(stream, events=("start", "end")):
            tag = node.tag
            if event == "start" and tag == f"{{{_MAIN_NS}}}row":
                current_row = int(node.attrib["r"])
                if current_row <= last_row:
                    raise CosmicDataError(f"Duplicate or out-of-order XML row identity in {sheet}")
                if node.get("hidden", "0") not in ("0", "false"):
                    raise CosmicDataError(f"Hidden source row is not in the reviewed schema: {sheet}!{current_row}")
                last_row, last_column = current_row, 0
            if event != "end":
                continue
            if tag == f"{{{_MAIN_NS}}}dimension":
                if dimension is not None:
                    raise CosmicDataError(f"Duplicate XML dimension in {sheet}")
                bounds = range_boundaries(node.attrib["ref"])
                if any(type(bound) is not int or bound < 1 for bound in bounds):
                    raise CosmicDataError(f"Invalid XML dimension in {sheet}")
                dimension = bounds
            elif tag == f"{{{_MAIN_NS}}}c":
                address = node.get("r", "")
                row, column = _coordinate(address, sheet)
                if row != current_row or column <= last_column:
                    raise CosmicDataError(f"Duplicate or out-of-order XML cell identity: {sheet}!{address}")
                if dimension is None or not (dimension[0] <= column <= dimension[2] and dimension[1] <= row <= dimension[3]):
                    raise CosmicDataError(f"XML cell outside declared dimension: {sheet}!{address}")
                if node.find(f"{{{_MAIN_NS}}}f") is not None:
                    raise CosmicDataError(f"Unexpected formula rather than source value: {sheet}!{address}")
                if node.get("t") == "e":
                    raise CosmicDataError(f"Excel error cell: {sheet}!{address}")
                last_column = column
                max_row, max_column = max(max_row, row), max(max_column, column)
                node.clear()
            elif tag == f"{{{_MAIN_NS}}}row":
                current_row = None
                node.clear()
            elif tag == f"{{{_MAIN_NS}}}col":
                if node.get("hidden", "0") not in ("0", "false"):
                    raise CosmicDataError(f"Hidden source column is not in the reviewed schema: {sheet}")
            elif tag == f"{{{_MAIN_NS}}}mergeCell":
                merges.append(node.attrib["ref"])
    if dimension is None or (max_row, max_column) != (dimension[3], dimension[2]):
        raise CosmicDataError(f"XML dimension does not match actual source-cell extent: {sheet}")
    _unique(merges, f"{sheet} merged range")
    if set(merges) != _expected_merges(sheet):
        raise CosmicDataError(f"Unexpected merged header groups in {sheet}")
    return (max_row, max_column), tuple(merges)


def _expect(sheet: CosmicSheet, row: int, column: int, expected: CellValue) -> None:
    cell = sheet.at(row, column)
    if cell.value != expected:
        raise CosmicDataError(f"{sheet.name}!{cell.coordinate}: expected {expected!r}, got {cell.value!r}")


def _expect_row(sheet: CosmicSheet, row: int, expected: Sequence[CellValue]) -> None:
    for column, value in enumerate(expected, 1):
        _expect(sheet, row, column, value)


def _number(cell: CosmicCell) -> int | float:
    if type(cell.value) not in (int, float) or not math.isfinite(cell.value):
        raise CosmicDataError(f"{cell.sheet}!{cell.coordinate}: expected a finite numeric source value, got {cell.value!r}")
    return cell.value


def _numeric_rectangle(sheet: CosmicSheet, first_row: int, first_column: int) -> None:
    for row in range(first_row, sheet.shape[0] + 1):
        for column in range(first_column, sheet.shape[1] + 1):
            _number(sheet.at(row, column))


def _validate_schema(data: CosmicDataset) -> None:
    index = data.sheets["Index"]
    _expect_row(index, 1, (None, None))
    for row, description in enumerate(_INDEX_DESCRIPTIONS, 2):
        _expect_row(index, row, (f"ST{row - 1}", description))
    design = data.sheets["ST1"]
    _expect_row(design, 1, (None, "Variable Levels", None, None))
    _expect_row(design, 2, ("Vessel", "O2", "AAs", "Glc"))
    _expect_row(design, 13, (None, None, None, None))
    _expect_row(design, 14, (_DESIGN_NOTE, None, None, None))
    for row in range(3, 13):
        _identifier(design.at(row, 1).value, f"ST1!A{row}")
        for column in range(2, 5):
            value = _number(design.at(row, column))
            if value not in (-1, 0, 1):
                raise CosmicDataError(f"ST1: unsupported coded variable level at row {row}")
    _unique(data.vessel_ids("ST1"), "ST1 vessel")
    process = data.sheets["ST2"]
    _expect_row(process, 1, _PROCESS_LABELS)
    _expect_row(process, 2, (None, "(day)", *["(dimensionless)"] * 26))
    _numeric_rectangle(process, 3, 2)
    identities = []
    last_time = {}
    for row in range(3, process.shape[0] + 1):
        vessel = _identifier(process.at(row, 1).value, f"ST2!A{row}")
        time = process.at(row, 2).value
        identities.append((vessel, time))
        if vessel in last_time and time < last_time[vessel]:
            raise CosmicDataError(f"ST2!B{row}: times must be increasing within each vessel")
        last_time[vessel] = time
    _unique(identities, "ST2 vessel/time")
    if len(last_time) != 10:
        raise CosmicDataError("ST2 must retain ten distinct source vessels")
    rates = data.sheets["ST3"]
    _expect_row(rates, 1, (None, None, "Growth Phase", *[None] * 9, "Production Phase", *[None] * 9))
    _expect(rates, 2, 1, None)
    _expect(rates, 2, 2, "Unit")
    for first in (3, 13):
        vessels = tuple(_identifier(rates.at(2, c).value, f"ST3!{get_column_letter(c)}2") for c in range(first, first + 10))
        _unique(vessels, "ST3 phase/vessel")
    for row, label in enumerate(_PROCESS_LABELS[3:], 3):
        unit = "1/day" if row <= 4 else "mg/10^6 cells-day" if row == 8 else "mmol/10^6 cells-day"
        _expect(rates, row, 1, label)
        _expect(rates, row, 2, unit)
    _numeric_rectangle(rates, 3, 3)
    objectives = data.sheets["ST4"]
    _expect(objectives, 1, 1, None)
    _expect(objectives, 2, 1, None)
    _expect(objectives, 3, 1, "Priority #")
    for row in range(4, 14):
        _expect(objectives, row, 1, row - 3)
    for first in range(2, 42, 4):
        _identifier(objectives.at(1, first).value, f"ST4!{get_column_letter(first)}1")
        for offset in range(1, 4):
            _expect(objectives, 1, first + offset, None)
        for offset, phase in ((0, _PHASES[0]), (2, _PHASES[1])):
            left = first + offset
            _expect(objectives, 2, left, phase)
            _expect(objectives, 2, left + 1, None)
            _expect(objectives, 3, left, "Reaction")
            _expect(objectives, 3, left + 1, "Efficiency")
            reactions = []
            missing_started = False
            for row in range(4, 14):
                reaction, efficiency = objectives.at(row, left), objectives.at(row, left + 1)
                if (reaction.value is None) != (efficiency.value is None):
                    raise CosmicDataError(f"ST4!{reaction.coordinate}:{efficiency.coordinate}: incomplete objective pair")
                if reaction.value is None:
                    missing_started = True
                else:
                    if missing_started:
                        raise CosmicDataError(f"ST4!{reaction.coordinate}: objective gap before populated priority")
                    reactions.append(_identifier(reaction.value, f"ST4!{reaction.coordinate}"))
                    _number(efficiency)
            if not reactions:
                raise CosmicDataError(f"ST4: entirely missing objective block at column {left}")
            _unique(reactions, f"ST4 objective reaction in column {left}")
    _unique(data.vessel_ids("ST4"), "ST4 vessel")
    fluxes = data.sheets["ST5"]
    _expect(fluxes, 1, 1, "Time (days) -->")
    times = tuple(_number(fluxes.at(1, c)) for c in range(2, fluxes.shape[1] + 1))
    _unique(times, "ST5 time")
    if any(left >= right for left, right in zip(times, times[1:])):
        raise CosmicDataError("ST5: time values must be strictly increasing")
    reactions = tuple(_identifier(row[0], f"ST5!A{r}") for r, row in enumerate(fluxes.values[1:], 2))
    _unique(reactions, "ST5 reaction")
    _numeric_rectangle(fluxes, 2, 2)
    if data.source.kind == "publisher_workbook":
        for sheet in data.sheets.values():
            if sheet.nonempty_cell_count != _NONEMPTY_CELLS[sheet.name]:
                raise CosmicDataError(f"Publisher nonempty-cell profile changed: {sheet.name}")
        for name in ("ST1", "ST2", "ST3", "ST4"):
            if data.vessel_ids(name) != _VESSELS:
                raise CosmicDataError(f"Publisher vessel identity profile changed: {name}")


def _parse(payload: bytes, source: CosmicSource) -> CosmicDataset:
    synthetic = source.kind == "synthetic_fixture"
    try:
        with ZipFile(BytesIO(payload)) as archive:
            members = _sheet_members(archive, synthetic)
            layouts = {sheet: _xml_layout(archive, member, sheet) for sheet, member in members.items()}
        for name, (shape, _) in layouts.items():
            if synthetic and name in ("ST2", "ST5"):
                valid = shape[0] >= (3 if name == "ST2" else 2) and (shape[1] == 28 if name == "ST2" else shape[1] >= 2)
            else:
                valid = shape == _SHAPES[name]
            if not valid:
                raise CosmicDataError(f"Unexpected {name} shape: {shape}")
        workbook = openpyxl.load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
        try:
            tables = {}
            for worksheet in workbook.worksheets:
                shape, merges = layouts[worksheet.title]
                rows = tuple(tuple(row) for row in worksheet.iter_rows(
                    min_row=1, min_col=1, max_row=shape[0], max_col=shape[1], values_only=True,
                ))
                if len(rows) != shape[0] or any(len(row) != shape[1] for row in rows):
                    raise CosmicDataError(f"Reader dropped source rows/columns in {worksheet.title}")
                for r, row in enumerate(rows, 1):
                    for c, value in enumerate(row, 1):
                        if type(value) not in (str, int, float, type(None)) or (type(value) is float and not math.isfinite(value)):
                            raise CosmicDataError(f"{worksheet.title}!{get_column_letter(c)}{r}: unsupported source cell value {value!r}")
                tables[worksheet.title] = CosmicSheet(source, worksheet.title, rows, merges)
        finally:
            workbook.close()
    except CosmicDataError:
        raise
    except (BadZipFile, ET.ParseError, KeyError, TypeError, ValueError, OSError) as exc:
        raise CosmicDataError(f"Invalid COSMIC workbook package: {exc}") from exc
    data = CosmicDataset(source, MappingProxyType(tables), source.sha256)
    _validate_schema(data)
    return data


def load_cosmic_workbook(workbook: str | Path) -> CosmicDataset:
    path = Path(workbook)
    payload = path.read_bytes()
    digest = hashlib.sha256(payload).hexdigest()
    if digest != PUBLISHER_WORKBOOK_SHA256 or len(payload) != PUBLISHER_WORKBOOK_BYTES:
        raise CosmicDataError("Publisher workbook SHA-256/byte-count mismatch; filenames do not establish source identity")
    source = CosmicSource("publisher_workbook", digest, len(payload), "10.1016/j.ymben.2024.02.012", "CHO (Chinese hamster ovary), not yeast")
    data = _parse(payload, source)
    after = path.read_bytes()
    after_digest = hashlib.sha256(after).hexdigest()
    if after_digest != digest or len(after) != len(payload):
        raise CosmicDataError("Source workbook changed during read-only inspection")
    return replace(data, sha256_after=after_digest)


def parse_synthetic_cosmic_workbook(payload: bytes) -> CosmicDataset:
    if not isinstance(payload, bytes):
        raise CosmicDataError("Synthetic parser requires explicit workbook bytes")
    digest = hashlib.sha256(payload).hexdigest()
    if digest == PUBLISHER_WORKBOOK_SHA256:
        raise CosmicDataError("Publisher bytes must use the pinned publisher loader, not the SYNTHETIC parser")
    return _parse(payload, CosmicSource("synthetic_fixture", digest, len(payload), None, "synthetic CHO-shaped fixture, not biological evidence"))


def _mapping_report(mapping: VesselMapping) -> dict:
    return {"method": mapping.method, "matched": list(mapping.matched), "workbook_only": list(mapping.workbook_only), "target_only": list(mapping.target_only)}


def _cell_report(cell: CosmicCell) -> dict:
    return {"source_sheet": cell.sheet, "source_cell": cell.coordinate, "value": cell.value}


def summarize_cosmic_data(data: CosmicDataset) -> dict:
    result = {
        "schema_version": 1,
        "source": asdict(data.source),
        "sha256_before": data.source.sha256,
        "sha256_after": data.sha256_after,
        "source_unchanged": data.source.sha256 == data.sha256_after,
        "exact_reproduction_ready": False,
        "independent_validation": False,
        "raw_data_publication_approved": False,
        "scope": "Source-workbook inspection, not model fitting, independent measured-flux truth, or yeast validation",
        "missing_reproduction_inputs": [
            "Original simulation-ready dFBA_data structure and its condition/component mappings",
            "Exact CHO metabolic/secretion model with stoichiometry, bounds and modifications",
            "Complete kinetic Vm_growth/Vm_prod/Km and phase-classification fitted coefficients",
            "Original physical initial/feed concentrations, normalization scales and perfusion configuration",
            "Complete original fitting/training setup and mapping of ST5 to a vessel/condition",
        ],
        "normalization": {
            "design_note": _cell_report(data.cell("ST1", "A14")),
            "process_description": _cell_report(data.cell("Index", "B3")),
            "physical_level_mapping_supplied": False,
            "temperature_factor_supplied": False,
            "process_denormalization_scales_supplied": False,
            "process_normalization_grouping_supplied": False,
            "production_phase_fraction_role": EvidenceRole.STATE_ANNOTATION.value,
            "production_phase_fraction_method_supplied_in_workbook": False,
            "transformations_applied": [],
        },
        "vessel_mappings": {
            f"ST1_to_{name}": _mapping_report(_match_vessels(data.vessel_ids("ST1"), data.vessel_ids(name)))
            for name in ("ST2", "ST3", "ST4")
        },
        "sheets": {},
    }
    result["vessel_mappings"]["ST5"] = {
        "status": "not_supplied", "vessel_id": None,
        "unmatched_workbook_vessels": list(data.vessel_ids("ST1")),
        "reason": "ST5 supplies reaction/time axes but no vessel/condition label; no ordinal join is applied",
    }
    result["vessel_mappings"]["external_reactor_names"] = {
        "status": "not_supplied", "mapping": [],
        "reason": "No differently named paper/code reactor-ID crosswalk is supplied in these sheets",
    }
    starts = {"Index": 2, "ST1": 3, "ST2": 3, "ST3": 3, "ST4": 4, "ST5": 2}
    for name, sheet in data.sheets.items():
        first = starts[name]
        last = 12 if name == "ST1" else sheet.shape[0]
        header_rows = range(1, first)
        summary = {
            "shape": list(sheet.shape), "nonempty_rows": sheet.nonempty_row_count,
            "nonempty_cells": sheet.nonempty_cell_count, "blank_rows": list(sheet.blank_rows),
            "data_rows": last - first + 1, "data_row_range": [first, last],
            "merged_ranges": list(sheet.merged_ranges), "hidden_rows": [], "hidden_columns": [],
            "formula_cells": [],
            "missing_cells": [f"{get_column_letter(c)}{r}" for r, row in enumerate(sheet.values, 1) for c, value in enumerate(row, 1) if value is None],
            "headers": {f"{get_column_letter(c)}{r}": value for r in header_rows for c, value in enumerate(sheet.values[r - 1], 1) if value is not None},
            "vessel_ids": list(data.vessel_ids(name)),
        }
        if name == "Index":
            summary["entries"] = [_cell_report(sheet.at(row, 2)) for row in range(2, 7)]
            counts = {}
        elif name == "ST1":
            counts = {EvidenceRole.SOURCE_DESIGN.value: 30}
            summary["coded_design"] = [
                {"vessel": _cell_report(sheet.at(row, 1)), "levels": [_cell_report(sheet.at(row, c)) for c in range(2, 5)]}
                for row in range(3, 13)
            ]
            summary["source_units"] = None
        elif name == "ST2":
            count = sheet.shape[0] - 2
            counts = {EvidenceRole.STATE_ANNOTATION.value: count, EvidenceRole.NORMALIZED_MEASUREMENT.value: 25 * count}
            summary["source_units"] = [_cell_report(sheet.at(2, c)) for c in range(2, 29)]
            summary["outside_unit_interval_cells"] = {
                role: [_cell_report(sheet.at(r, c)) for r in range(3, sheet.shape[0] + 1)
                       for c in columns if not 0 <= sheet.at(r, c).value <= 1]
                for role, columns in (("state_fraction", (3,)), ("normalized_process", range(4, 29)))
            }
            summary["vessel_observations"] = [
                {"vessel": vessel, "rows": [r for r, row in enumerate(sheet.values, 1) if row[0] == vessel],
                 "times": [_cell_report(sheet.at(r, 2)) for r, row in enumerate(sheet.values, 1) if row[0] == vessel]}
                for vessel in data.vessel_ids(name)
            ]
        elif name == "ST3":
            counts = {EvidenceRole.COMPUTED_RATE.value: 500}
            summary["source_units"] = [
                {"quantity": _cell_report(sheet.at(r, 1)), "unit": _cell_report(sheet.at(r, 2))} for r in range(3, 28)
            ]
        elif name == "ST4":
            counts = {EvidenceRole.FITTED_OBJECTIVE.value: 200}
            summary["source_units"] = None
            summary["unit_status"] = "Efficiency has no explicit unit cell in the source"
            summary["objective_blocks"] = [
                {"vessel": _cell_report(sheet.at(1, 2 + 4 * ((c - 2) // 4))),
                 "phase": _cell_report(sheet.at(2, c - 1)), "efficiency_column": get_column_letter(c),
                 "populated_priorities": [r - 3 for r in range(4, 14) if sheet.at(r, c).value is not None],
                 "missing_priorities": [r - 3 for r in range(4, 14) if sheet.at(r, c).value is None]}
                for c in range(3, 42, 2)
            ]
            efficiencies = [row[c - 1] for row in sheet.values[3:] for c in range(3, 42, 2) if row[c - 1] is not None]
            summary["efficiencies_above_one_preserved"] = sum(value > 1 for value in efficiencies)
            summary["zero_efficiencies_preserved"] = sum(value == 0 for value in efficiencies)
        else:
            counts = {EvidenceRole.COMPUTED_FLUX.value: (sheet.shape[0] - 1) * (sheet.shape[1] - 1)}
            summary["source_units"] = {"value": "mmol/gDW-d", "source_sheet": "Index", "source_cell": "B6"}
            summary["source_time_unit"] = {"value": "days", "source_sheet": "ST5", "source_cell": "A1"}
            summary["reaction_count"] = sheet.shape[0] - 1
            summary["time_count"] = sheet.shape[1] - 1
            summary["first_reaction"] = _cell_report(sheet.at(2, 1))
            summary["last_reaction"] = _cell_report(sheet.at(sheet.shape[0], 1))
            summary["time_axis"] = [_cell_report(sheet.at(1, c)) for c in range(2, sheet.shape[1] + 1)]
            signs = Counter("zero" if value == 0 else "negative" if value < 0 else "positive" for row in sheet.values[1:] for value in row[1:])
            summary.update({f"{sign}_fluxes": signs[sign] for sign in ("zero", "negative", "positive")})
        summary["quantities_by_evidence_role"] = counts
        summary["quantity_slots"] = sum(counts.values())
        summary["missing_quantity_slots"] = sum(q.is_missing for q in data.iter_quantities(name)) if name == "ST4" else 0
        result["sheets"][name] = summary
    objective_ids = {q.label.value for q in data.iter_quantities("ST4") if not q.is_missing}
    flux_ids = {row[0] for row in data.sheets["ST5"].values[1:]}
    result["objective_reaction_links_to_ST5"] = {
        "method": "exact_reaction_label_only_not_vessel_or_model_equivalence",
        "matched": sorted(objective_ids & flux_ids), "objective_only": sorted(objective_ids - flux_ids),
        "flux_only_count": len(flux_ids - objective_ids),
    }
    return result
