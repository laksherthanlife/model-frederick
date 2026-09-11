"""Reader for the per-construct dose-response sheets.

Each biosensor plate carries a sheet per construct (UPRE1, UPRE2, NativeYap1,
AlteredYap1) holding normalised fluorescence against time, one column per DTT dose:

    Time | 0 mM | 0.1 mM | 0.2 mM | 0.5 mM | 1 mM | 2 mM | 5 mM

This is the reporter response as the team actually uses it, which is what G4 has to
be run against -- not a placeholder, and not a quantity only this codebase computes.
"""

import numpy as np
import pandas as pd
import pytest

from ystwin.plate.dose_response import read_dose_response


@pytest.fixture
def dose_sheet(tmp_path):
    def _make(name="plate.xlsx", lead_rows=3, constructs=("UPRE1", "NativeYap1"),
              with_raw_sheets=False):
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        if with_raw_sheets:
            od = wb.create_sheet("OD600")
            od.append(["Time", "A1", "A2"])
            od.append(["00:07:25", 0.11, 0.12])
        for k, construct in enumerate(constructs):
            ws = wb.create_sheet(construct)
            for _ in range(lead_rows):
                ws.append([])
            ws.append([None, "Time", "0 mM", "0.2 mM", "0.5 mM"])
            for i, t in enumerate([0, 10, 20, 30]):
                base = 3000 + 100 * k
                ws.append([None, t, base + i, base + 2 * i + 50, base + 4 * i + 120])
        path = tmp_path / name
        wb.save(path)
        return path

    return _make


def test_one_row_per_construct_dose_and_time(dose_sheet):
    tidy = read_dose_response(dose_sheet())

    assert set(tidy.construct) == {"UPRE1", "NativeYap1"}
    assert sorted(tidy.dose_mM.unique()) == [0.0, 0.2, 0.5]
    assert len(tidy) == 2 * 3 * 4


def test_time_is_converted_from_minutes_to_hours(dose_sheet):
    tidy = read_dose_response(dose_sheet())

    assert sorted(tidy.time_h.unique()) == pytest.approx([0.0, 1 / 6, 2 / 6, 3 / 6])


def test_the_signal_values_survive(dose_sheet):
    tidy = read_dose_response(dose_sheet())
    row = tidy.query("construct == 'UPRE1' and dose_mM == 0.5 and time_h == 0.0")

    assert row.signal.iloc[0] == pytest.approx(3120)


def test_a_varying_preamble_does_not_shift_the_parse(dose_sheet):
    a = read_dose_response(dose_sheet(lead_rows=1))
    b = read_dose_response(dose_sheet(name="b.xlsx", lead_rows=5))

    pd.testing.assert_frame_equal(
        a.sort_values(list(a.columns)).reset_index(drop=True),
        b.sort_values(list(b.columns)).reset_index(drop=True),
    )


def test_sheets_that_are_not_construct_names_are_ignored(dose_sheet):
    path = dose_sheet(with_raw_sheets=True)

    assert set(read_dose_response(path).construct) == {"UPRE1", "NativeYap1"}


def test_a_workbook_with_no_construct_sheet_is_refused(tmp_path):
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "OD600"
    wb.active.append(["Time", "A1"])
    path = tmp_path / "none.xlsx"
    wb.save(path)

    with pytest.raises(ValueError, match="no per-construct"):
        read_dose_response(path)


def test_the_endpoint_response_summarises_each_dose(dose_sheet):
    from ystwin.plate.dose_response import endpoint_response

    tidy = read_dose_response(dose_sheet())
    summary = endpoint_response(tidy, fraction=0.5)

    row = summary.query("construct == 'UPRE1' and dose_mM == 0.0")
    assert row.response.iloc[0] == pytest.approx(np.mean([3002, 3003]))


def test_the_endpoint_response_is_expressed_relative_to_the_untreated_dose(dose_sheet):
    from ystwin.plate.dose_response import endpoint_response

    tidy = read_dose_response(dose_sheet())
    summary = endpoint_response(tidy, fraction=0.5, relative_to_dose=0.0)

    control = summary.query("construct == 'UPRE1' and dose_mM == 0.0")
    assert control.response.iloc[0] == pytest.approx(1.0)
