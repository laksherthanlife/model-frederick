from __future__ import annotations

from collections import Counter
from copy import deepcopy
import hashlib
from io import BytesIO
import json
from pathlib import Path
import runpy
import subprocess
import sys
import xml.etree.ElementTree as ET
from zipfile import ZipFile

import openpyxl
import pytest

from ystwin.analysis.cosmic_data import (
    CosmicDataError,
    EvidenceRole,
    PUBLISHER_WORKBOOK_BYTES,
    PUBLISHER_WORKBOOK_SHA256,
    QuantityRole,
    SYNTHETIC_WORKBOOK_TITLE,
    WORKBOOK_FILENAMES,
    load_cosmic_workbook,
    parse_synthetic_cosmic_workbook,
    summarize_cosmic_data,
)


ROOT = Path(__file__).resolve().parents[1]
SCRIPT = ROOT / "scripts/inspect_cosmic_data.py"
SYNTHETIC_VESSELS = tuple(f"SYNTHETIC_V{i:02d}" for i in range(1, 11))
PROCESS_LABELS = (
    "Vessel", "Time", "Production phase fraction", "Cell Density", "Cell Volume",
    "Glucose", "Lactate", "NH4", "Titer", "Glutamine", "Glutamate", "L-Asparagine",
    "L-Aspartic acid", "L-Serine", "Glycine", "L-Alanine ", "L-Proline", "L-Threonine",
    "L-Histidine", "L-Lysine", "L-Valine ", "L-Methionine", "L-Arginine", "L-Tyrosine",
    "L-Isoleucine", "L-Leucine", "L-Phenylalanine", "L-Tryptophan",
)
INDEX_DESCRIPTIONS = (
    "Box Behnken design with 3 factors",
    "Normalized measured process data",
    "Computed uptake and secretion rates",
    "Cellular objectives in the growth and production phases",
    "Intracellular fluxes (mmol/gDW-d) computed using COSMIC-dFBA",
)


def _synthetic_workbook():
    workbook = openpyxl.Workbook()
    workbook.properties.title = SYNTHETIC_WORKBOOK_TITLE
    workbook.properties.description = "Invented parser values; not publisher observations or outputs"
    index = workbook.active
    index.title = "Index"
    for row, description in enumerate(INDEX_DESCRIPTIONS, 2):
        index.cell(row, 1, f"ST{row - 1}")
        index.cell(row, 2, description)
    design = workbook.create_sheet("ST1")
    design["B1"] = "Variable Levels"
    design.merge_cells("B1:D1")
    design.append(["Vessel", "O2", "AAs", "Glc"])
    for i, vessel in enumerate(SYNTHETIC_VESSELS):
        design.append([vessel, i % 3 - 1, (i + 1) % 3 - 1, (i + 2) % 3 - 1])
    design["A14"] = (
        "*0 represents a central value, +1 and -1 represent increased and reduced values "
        "for the independent variable, respectively"
    )
    process = workbook.create_sheet("ST2")
    process.append(PROCESS_LABELS)
    process.append([None, "(day)", *["(dimensionless)"] * 26])
    for vessel in SYNTHETIC_VESSELS:
        process.append([vessel, 0, 0, *[0.25] * 25])
        process.append([vessel, 2.5, 1, *[0.75] * 25])
    rates = workbook.create_sheet("ST3")
    rates["C1"] = "Growth Phase"
    rates["M1"] = "Production Phase"
    rates.merge_cells("C1:L1")
    rates.merge_cells("M1:V1")
    rates.append([None, "Unit", *SYNTHETIC_VESSELS, *SYNTHETIC_VESSELS])
    for i, label in enumerate(PROCESS_LABELS[3:]):
        unit = "1/day" if i < 2 else "mg/10^6 cells-day" if label == "Titer" else "mmol/10^6 cells-day"
        rates.append([label, unit, *[(-1 if i % 2 else 1) * (i + 1.25)] * 20])
    objectives = workbook.create_sheet("ST4")
    objectives["A3"] = "Priority #"
    for priority in range(1, 11):
        objectives.cell(priority + 3, 1, priority)
    for i, vessel in enumerate(SYNTHETIC_VESSELS):
        column = 2 + 4 * i
        objectives.cell(1, column, vessel)
        objectives.merge_cells(start_row=1, end_row=1, start_column=column, end_column=column + 3)
        for offset, phase in ((0, "Growth Phase"), (2, "Production Phase")):
            left = column + offset
            objectives.cell(2, left, phase)
            objectives.merge_cells(start_row=2, end_row=2, start_column=left, end_column=left + 1)
            objectives.cell(3, left, "Reaction")
            objectives.cell(3, left + 1, "Efficiency")
            objectives.cell(4, left, "SYNTHETIC_RXN_1")
            objectives.cell(4, left + 1, 1.00000000005)
            objectives.cell(5, left, "SYNTHETIC_RXN_2")
            objectives.cell(5, left + 1, 0)
    fluxes = workbook.create_sheet("ST5")
    fluxes.append(["Time (days) -->", 0.125, 0.75, 2.5])
    fluxes.append(["SYNTHETIC_RXN_1", 0, -1e-20, 2])
    fluxes.append(["SYNTHETIC_RXN_2", -2, 0, 1e-15])
    fluxes.append(["SYNTHETIC_RXN_3", 3, -3, 0])
    return workbook


def _bytes(workbook):
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.fixture
def synthetic_payload():
    return _bytes(_synthetic_workbook())


@pytest.fixture
def synthetic_data(synthetic_payload):
    return parse_synthetic_cosmic_workbook(synthetic_payload)


def test_synthetic_fixture_cannot_inherit_publisher_identity(synthetic_payload, synthetic_data):
    assert PUBLISHER_WORKBOOK_SHA256 == "a68dc8bf86b87e6b9d4f732cd3413948bdb567bf3f85fdabe937a202bb7a7daa"
    assert PUBLISHER_WORKBOOK_BYTES == 4156209
    assert synthetic_data.source.kind == "synthetic_fixture"
    assert synthetic_data.source.doi is None
    assert synthetic_data.source.sha256 == hashlib.sha256(synthetic_payload).hexdigest()
    assert synthetic_data.source.sha256 != PUBLISHER_WORKBOOK_SHA256
    assert synthetic_data.source.byte_count == len(synthetic_payload)
    assert synthetic_data.sha256_after == synthetic_data.source.sha256
    assert synthetic_data.cell("ST5", "D4").source.kind == "synthetic_fixture"


def test_all_raw_rows_columns_blanks_and_merged_anchors_survive(synthetic_data):
    assert tuple(synthetic_data.sheets) == ("Index", "ST1", "ST2", "ST3", "ST4", "ST5")
    design = synthetic_data.sheets["ST1"]
    assert design.shape == (14, 4)
    assert design.nonempty_row_count == 13
    assert design.blank_rows == (13,)
    assert design.merged_ranges == ("B1:D1",)
    assert design.cell("B1").value == "Variable Levels"
    assert design.cell("C1").value is None
    assert design.cell("D13").value is None
    frame = design.to_frame()
    assert list(frame.index) == list(range(1, 15))
    assert list(frame.columns) == ["A", "B", "C", "D"]
    assert frame.at[13, "D"] is None
    assert frame.at[3, "B"] == -1
    frame.at[3, "B"] = 999
    assert design.cell("B3").value == -1
    with pytest.raises(TypeError):
        synthetic_data.sheets["ST1"] = None
    assert synthetic_data.cell("ST2", "P1").value == "L-Alanine "
    assert synthetic_data.cell("ST3", "A20").value == "L-Valine "


def test_quantity_identity_and_metadata_are_original_cell_references(synthetic_data):
    quantity = synthetic_data.quantity("ST2", "D3")
    assert quantity.cell.coordinate == "D3"
    assert quantity.cell.sheet == "ST2"
    assert quantity.cell.identity == f"synthetic_fixture:{synthetic_data.source.sha256}:ST2!D3"
    assert quantity.label.coordinate == "D1"
    assert quantity.label.value == "Cell Density"
    assert quantity.source_unit == "(dimensionless)"
    assert quantity.unit_cell.coordinate == "D2"
    assert quantity.vessel.coordinate == "A3"
    assert quantity.vessel.value == "SYNTHETIC_V01"
    assert quantity.time.coordinate == "B3"
    assert quantity.time.value == 0
    assert quantity.source_time_unit == "(day)"
    assert quantity.time_unit_cell.coordinate == "B2"
    assert quantity.state_fraction.coordinate == "C3"
    assert quantity.role_evidence.sheet == "Index"
    assert quantity.role_evidence.coordinate == "B3"
    assert quantity.quantity_role == QuantityRole.NORMALIZED_PROCESS_QUANTITY
    assert quantity.evidence_role == EvidenceRole.NORMALIZED_MEASUREMENT
    assert quantity.normalization == "source_normalized_scale_not_supplied"
    assert quantity.value == 0.25


def test_state_annotations_are_not_measured_process_quantities(synthetic_data):
    phase = synthetic_data.quantity("ST2", "C3")
    assert phase.quantity_role == QuantityRole.PRODUCTION_PHASE_FRACTION
    assert phase.evidence_role == EvidenceRole.STATE_ANNOTATION
    assert phase.evidence_role != EvidenceRole.NORMALIZED_MEASUREMENT
    assert phase.source_unit == "(dimensionless)"
    coded = synthetic_data.quantity("ST1", "B3")
    assert coded.quantity_role == QuantityRole.CODED_DESIGN_LEVEL
    assert coded.source_unit is None and coded.unit_cell is None
    assert coded.normalization_cell.coordinate == "A14"
    assert coded.vessel.value == "SYNTHETIC_V01"


def test_computed_outputs_keep_source_units_and_no_unjustified_vessel_join(synthetic_data):
    density = synthetic_data.quantity("ST3", "C3")
    metabolite = synthetic_data.quantity("ST3", "M5")
    titer = synthetic_data.quantity("ST3", "V8")
    assert density.source_unit == "1/day"
    assert metabolite.source_unit == "mmol/10^6 cells-day"
    assert metabolite.unit_cell.coordinate == "B5"
    assert metabolite.phase.coordinate == "M1"
    assert metabolite.phase.value == "Production Phase"
    assert metabolite.vessel.coordinate == "M2"
    assert titer.source_unit == "mg/10^6 cells-day"
    assert density.evidence_role == metabolite.evidence_role == EvidenceRole.COMPUTED_RATE
    flux = synthetic_data.quantity("ST5", "C2")
    assert flux.quantity_role == QuantityRole.INTRACELLULAR_FLUX
    assert flux.evidence_role == EvidenceRole.COMPUTED_FLUX
    assert flux.value == -1e-20
    assert flux.label.value == "SYNTHETIC_RXN_1"
    assert flux.label.coordinate == "A2"
    assert flux.source_unit == "mmol/gDW-d"
    assert (flux.unit_cell.sheet, flux.unit_cell.coordinate) == ("Index", "B6")
    assert flux.time.coordinate == "C1" and flux.time.value == 0.75
    assert flux.source_time_unit == "days"
    assert flux.vessel is None and flux.phase is None


def test_source_annotations_and_normalized_values_are_not_clamped_or_range_filtered():
    workbook = _synthetic_workbook()
    workbook["ST2"]["C3"] = -1e-17
    workbook["ST2"]["L3"] = -0.125
    workbook["ST2"]["M3"] = -0.25
    data = parse_synthetic_cosmic_workbook(_bytes(workbook))
    assert data.quantity("ST2", "C3").value == -1e-17
    assert data.quantity("ST2", "L3").value == -0.125
    assert data.quantity("ST2", "M3").value == -0.25
    assert len(list(data.iter_quantities("ST2"))) == 520
    excursions = summarize_cosmic_data(data)["sheets"]["ST2"]["outside_unit_interval_cells"]
    assert [cell["source_cell"] for cell in excursions["state_fraction"]] == ["C3"]
    assert [cell["source_cell"] for cell in excursions["normalized_process"]] == ["L3", "M3"]


def test_objective_slots_preserve_missing_pairs_zero_and_above_one(synthetic_data):
    quantities = list(synthetic_data.iter_quantities("ST4"))
    assert len(quantities) == 200
    assert sum(quantity.is_missing for quantity in quantities) == 160
    present = synthetic_data.quantity("ST4", "C4")
    assert present.evidence_role == EvidenceRole.FITTED_OBJECTIVE
    assert present.value == 1.00000000005
    assert present.source_unit is None and present.unit_cell is None
    assert present.label.coordinate == "B4"
    assert present.vessel.coordinate == "B1"
    assert present.phase.coordinate == "B2"
    assert present.priority.coordinate == "A4" and present.priority.value == 1
    assert synthetic_data.quantity("ST4", "C5").value == 0
    assert not synthetic_data.quantity("ST4", "C5").is_missing
    missing = synthetic_data.quantity("ST4", "AO13")
    assert missing.value is None and missing.label.value is None
    assert missing.label.coordinate == "AN13"
    assert missing.priority.value == 10
    assert missing.vessel.coordinate == "AL1"
    assert missing.phase.coordinate == "AN2"


def test_quantity_enumeration_never_filters_zero_missing_or_tiny_values(synthetic_data):
    expected_counts = {"ST1": 30, "ST2": 520, "ST3": 500, "ST4": 200, "ST5": 9}
    identities = []
    for sheet, count in expected_counts.items():
        quantities = list(synthetic_data.iter_quantities(sheet))
        assert len(quantities) == count
        identities.extend(quantity.cell.identity for quantity in quantities)
    assert len(identities) == len(set(identities))
    flux_values = [quantity.value for quantity in synthetic_data.iter_quantities("ST5")]
    assert flux_values == [0, -1e-20, 2, -2, 0, 1e-15, 3, -3, 0]


@pytest.mark.parametrize("sheet,coordinate", [("ST1", "A13"), ("ST2", "B3"), ("ST4", "B4"), ("Index", "B3")])
def test_headers_axes_and_blank_separators_cannot_masquerade_as_quantities(synthetic_data, sheet, coordinate):
    with pytest.raises(CosmicDataError, match="quantity cell"):
        synthetic_data.quantity(sheet, coordinate)


@pytest.mark.parametrize("coordinate", ["A0", "A01", "a1", "A15", "E1", "A1:B2"])
def test_cell_access_is_bounded_and_canonical(synthetic_data, coordinate):
    with pytest.raises(CosmicDataError, match="cell"):
        synthetic_data.cell("ST1", coordinate)


def test_vessel_matching_is_exact_and_reports_unmatched_names(synthetic_data):
    assert synthetic_data.vessel_ids("ST2") == SYNTHETIC_VESSELS
    assert synthetic_data.vessel_ids("ST5") == ()
    match = synthetic_data.match_vessels(["SYNTHETIC_V01", "BR01"])
    assert match.matched == ("SYNTHETIC_V01",)
    assert match.workbook_only == SYNTHETIC_VESSELS[1:]
    assert match.target_only == ("BR01",)
    with pytest.raises(CosmicDataError, match="Duplicate"):
        synthetic_data.match_vessels(["BR01", "BR01"])


def test_cross_sheet_names_are_not_relabelled_by_position():
    workbook = _synthetic_workbook()
    workbook["ST2"]["A3"] = "OTHER_VESSEL"
    workbook["ST2"]["A4"] = "OTHER_VESSEL"
    data = parse_synthetic_cosmic_workbook(_bytes(workbook))
    report = summarize_cosmic_data(data)
    mapping = report["vessel_mappings"]["ST1_to_ST2"]
    assert mapping["workbook_only"] == ["SYNTHETIC_V01"]
    assert mapping["target_only"] == ["OTHER_VESSEL"]
    assert data.quantity("ST2", "D3").vessel.value == "OTHER_VESSEL"


def test_summary_is_json_serializable_and_does_not_claim_reproduction(synthetic_data):
    report = summarize_cosmic_data(synthetic_data)
    json.dumps(report, allow_nan=False)
    assert report["source"]["kind"] == "synthetic_fixture"
    assert report["exact_reproduction_ready"] is False
    assert report["independent_validation"] is False
    assert report["raw_data_publication_approved"] is False
    assert report["sheets"]["ST2"]["data_rows"] == 20
    assert report["sheets"]["ST4"]["quantity_slots"] == 200
    assert report["sheets"]["ST4"]["missing_quantity_slots"] == 160
    assert report["sheets"]["ST5"]["shape"] == [4, 4]
    assert report["sheets"]["ST5"]["zero_fluxes"] == 3
    assert report["vessel_mappings"]["ST5"]["status"] == "not_supplied"
    assert report["normalization"]["physical_level_mapping_supplied"] is False
    assert report["normalization"]["process_denormalization_scales_supplied"] is False
    roles = report["sheets"]["ST2"]["quantities_by_evidence_role"]
    assert roles[EvidenceRole.NORMALIZED_MEASUREMENT.value] == 500
    assert roles[EvidenceRole.STATE_ANNOTATION.value] == 20


@pytest.mark.parametrize("filename", [*WORKBOOK_FILENAMES, "renamed.xlsx"])
def test_publisher_loader_requires_pinned_bytes_not_a_filename(tmp_path, synthetic_payload, filename):
    path = tmp_path / filename
    path.write_bytes(synthetic_payload)
    before = path.read_bytes()
    with pytest.raises(CosmicDataError, match="SHA-256"):
        load_cosmic_workbook(path)
    assert path.read_bytes() == before


def test_publisher_loader_rejects_html_and_has_no_identity_override(tmp_path):
    path = tmp_path / WORKBOOK_FILENAMES[0]
    path.write_bytes(b"<html>Not a scientific asset</html>")
    with pytest.raises(CosmicDataError, match="SHA-256"):
        load_cosmic_workbook(path)
    with pytest.raises(TypeError):
        load_cosmic_workbook(path, expected_sha256=hashlib.sha256(path.read_bytes()).hexdigest())
    with pytest.raises(TypeError):
        load_cosmic_workbook()


def test_synthetic_parser_requires_an_explicit_marker():
    workbook = _synthetic_workbook()
    workbook.properties.title = "An unreviewed workbook"
    with pytest.raises(CosmicDataError, match="SYNTHETIC"):
        parse_synthetic_cosmic_workbook(_bytes(workbook))


@pytest.mark.parametrize(
    "sheet,coordinate,value,match",
    [
        ("Index", "B3", "Raw physical concentrations", "Index!B3"),
        ("ST1", "B1", "Concentrations", "ST1!B1"),
        ("ST1", "B3", 2, "coded"),
        ("ST1", "A4", "SYNTHETIC_V01", "Duplicate"),
        ("ST2", "P1", "L-Alanine", "ST2!P1"),
        ("ST2", "D2", "mmol/L", "ST2!D2"),
        ("ST2", "B4", 0, "Duplicate"),
        ("ST2", "D3", None, "ST2!D3"),
        ("ST2", "D3", True, "ST2!D3"),
        ("ST2", "D3", "=1+2", "formula"),
        ("ST3", "C3", None, "ST3!C3"),
        ("ST3", "B8", "mmol/gDW/h", "ST3!B8"),
        ("ST3", "D2", "SYNTHETIC_V01", "Duplicate"),
        ("ST4", "F1", "SYNTHETIC_V01", "Duplicate"),
        ("ST4", "A5", 1, "ST4!A5"),
        ("ST4", "B5", "SYNTHETIC_RXN_1", "Duplicate"),
        ("ST4", "B5", None, "pair"),
        ("ST4", "A14", 11, "shape"),
        ("ST5", "A3", "SYNTHETIC_RXN_1", "Duplicate"),
        ("ST5", "C1", 0.125, "Duplicate"),
        ("ST5", "C1", 0.01, "increasing"),
        ("ST5", "B2", None, "ST5!B2"),
        ("ST5", "B2", "0.0", "ST5!B2"),
        ("ST5", "B2", "#DIV/0!", "error"),
        ("ST5", "E2", 42, "ST5!E1"),
    ],
)
def test_schema_drift_and_ambiguous_identities_fail_closed(sheet, coordinate, value, match):
    workbook = _synthetic_workbook()
    workbook[sheet][coordinate] = value
    with pytest.raises(CosmicDataError, match=match):
        parse_synthetic_cosmic_workbook(_bytes(workbook))


@pytest.mark.parametrize("mutation", ["row", "column", "sheet", "merge", "extra_sheet", "objective_gap"])
def test_hidden_data_and_layout_changes_cannot_be_silently_dropped(mutation):
    workbook = _synthetic_workbook()
    if mutation == "row":
        workbook["ST2"].row_dimensions[3].hidden = True
    elif mutation == "column":
        workbook["ST5"].column_dimensions["C"].hidden = True
    elif mutation == "sheet":
        workbook["ST5"].sheet_state = "hidden"
    elif mutation == "merge":
        workbook["ST3"].unmerge_cells("C1:L1")
    elif mutation == "extra_sheet":
        workbook.create_sheet("unreviewed")
    else:
        workbook["ST4"]["B4"] = None
        workbook["ST4"]["C4"] = None
    with pytest.raises(CosmicDataError):
        parse_synthetic_cosmic_workbook(_bytes(workbook))


def _rewrite_xml(payload, member, change):
    output = BytesIO()
    with ZipFile(BytesIO(payload)) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            body = source.read(info)
            if info.filename == member:
                root = ET.fromstring(body)
                change(root)
                body = ET.tostring(root)
            target.writestr(info, body)
    return output.getvalue()


@pytest.mark.parametrize("mutation", ["duplicate_cell", "duplicate_row", "truncated_dimension"])
def test_xml_identities_and_actual_extents_are_checked_before_openpyxl_drops_them(synthetic_payload, mutation):
    ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}

    def change(root):
        rows = root.find("s:sheetData", ns)
        if mutation == "duplicate_cell":
            rows[1].append(deepcopy(rows[1][1]))
        elif mutation == "duplicate_row":
            rows.append(deepcopy(rows[1]))
        else:
            root.find("s:dimension", ns).set("ref", "A1:B2")

    changed = _rewrite_xml(synthetic_payload, "xl/worksheets/sheet6.xml", change)
    with pytest.raises(CosmicDataError, match="Duplicate|dimension|order"):
        parse_synthetic_cosmic_workbook(changed)


def test_duplicate_zip_members_cannot_shadow_a_sheet(synthetic_payload):
    output = BytesIO()
    with ZipFile(BytesIO(synthetic_payload)) as source, ZipFile(output, "w") as target:
        for info in source.infolist():
            target.writestr(info, source.read(info))
        with pytest.warns(UserWarning, match="Duplicate name"):
            target.writestr("xl/worksheets/sheet6.xml", source.read("xl/worksheets/sheet6.xml"))
    with pytest.raises(CosmicDataError, match="Duplicate ZIP"):
        parse_synthetic_cosmic_workbook(output.getvalue())


def test_cli_loads_the_checkout_without_pythonpath_or_an_updated_install(tmp_path):
    result = subprocess.run([sys.executable, "-I", "-B", str(SCRIPT), "--help"],
                            cwd=tmp_path, capture_output=True, text=True)
    assert result.returncode == 0, result.stderr
    assert "--workbook" in result.stdout and "--report" in result.stdout
    assert not list(tmp_path.iterdir())


def test_cli_requires_explicit_input_and_never_uses_a_default_directory(tmp_path):
    result = subprocess.run([sys.executable, "-B", str(SCRIPT)], cwd=tmp_path, capture_output=True, text=True)
    assert result.returncode == 2
    assert "--input-dir" in result.stderr and "--workbook" in result.stderr
    assert not list(tmp_path.iterdir())


def test_cli_rejects_ambiguous_directory_instead_of_selecting_an_ordinal(tmp_path, synthetic_payload):
    for name in (*WORKBOOK_FILENAMES, "other.xlsx"):
        (tmp_path / name).write_bytes(synthetic_payload)
    result = subprocess.run(
        [sys.executable, "-B", str(SCRIPT), "--input-dir", str(tmp_path)], capture_output=True, text=True,
    )
    assert result.returncode == 2
    assert "ambiguous" in result.stderr.lower()
    assert len(list(tmp_path.iterdir())) == 3


def test_cli_report_guards_are_external_explicit_and_non_overwriting(tmp_path):
    cli = runpy.run_path(str(SCRIPT))
    validate = cli["_report_destination"]
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    source = source_dir / WORKBOOK_FILENAMES[0]
    source.write_bytes(b"synthetic guard input, never parsed")
    report = tmp_path / "inspection.json"
    assert validate(report, [source]) == report.resolve()
    for invalid in (Path("relative.json"), ROOT / "cosmic-test-report.json", source_dir / "report.json", source):
        with pytest.raises(CosmicDataError):
            validate(invalid, [source])
    alternate_root = ROOT.parent / ROOT.name.swapcase()
    if alternate_root.exists() and alternate_root.samefile(ROOT):
        with pytest.raises(CosmicDataError, match="outside the repository"):
            validate(alternate_root / "cosmic-test-report.json", [source])
    alternate_source = source_dir.parent / source_dir.name.swapcase()
    if alternate_source.exists() and alternate_source.samefile(source_dir):
        with pytest.raises(CosmicDataError, match="outside the source directory"):
            validate(alternate_source / "report.json", [source])
    report.write_text("do not overwrite", encoding="utf-8")
    with pytest.raises(CosmicDataError, match="exists"):
        validate(report, [source])
    assert report.read_text(encoding="utf-8") == "do not overwrite"
    assert source.read_bytes() == b"synthetic guard input, never parsed"
    link = tmp_path / "repo-link"
    link.symlink_to(ROOT, target_is_directory=True)
    with pytest.raises(CosmicDataError):
        validate(link / "cosmic-test-report.json", [source])


def test_cli_does_not_write_a_report_when_source_verification_fails(tmp_path, synthetic_payload):
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    source = source_dir / WORKBOOK_FILENAMES[0]
    source.write_bytes(synthetic_payload)
    report = tmp_path / "inspection.json"
    result = subprocess.run(
        [sys.executable, "-B", str(SCRIPT), "--workbook", str(source), "--report", str(report)],
        capture_output=True, text=True,
    )
    assert result.returncode == 2
    assert "SHA-256" in result.stderr
    assert not report.exists()
    assert source.read_bytes() == synthetic_payload


def test_report_role_counts_equal_quantity_iteration(synthetic_data):
    report = summarize_cosmic_data(synthetic_data)
    for sheet in ("ST1", "ST2", "ST3", "ST4", "ST5"):
        counts = Counter(item.evidence_role.value for item in synthetic_data.iter_quantities(sheet))
        assert report["sheets"][sheet]["quantities_by_evidence_role"] == dict(counts)
