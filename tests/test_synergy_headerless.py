"""Later exports drop the temperature column, leaving the sheet name as the only channel id.

    Time | A1 | A2 | B1 ...        rather than        Time | T° OD600:600 | A1 | ...

The block is still a kinetic block and must still be read; the channel name simply
has to come from somewhere else. Guessing from column position would be fragile, so
the sheet name is used, normalised.
"""

import pytest

from ystwin.plate.synergy import read_synergy_kinetic


def test_a_block_without_a_channel_column_is_still_read(make_headerless_channel_file):
    run = read_synergy_kinetic(make_headerless_channel_file())

    assert set(run.channel_names) == {"OD600", "mCitrine"}


def test_the_values_survive_the_missing_column(make_headerless_channel_file):
    run = read_synergy_kinetic(make_headerless_channel_file())

    assert run.channel("OD600").data["A1"].to_numpy() == pytest.approx([0.11, 0.14, 0.22])
    assert run.channel("mCitrine").data["B1"].to_numpy() == pytest.approx([340, 410, 630])


def test_sheet_names_are_normalised_to_channel_names(make_headerless_channel_file):
    run = read_synergy_kinetic(
        make_headerless_channel_file(sheet_names=("OD600 (Raw)", "mCitrine (Raw)"))
    )

    assert set(run.channel_names) == {"OD600", "mCitrine"}


def test_alternate_optical_density_sheet_spellings_resolve_to_one_channel(make_headerless_channel_file):
    for spelling in ("Raw OD", "NEW OD", "OD600", "OD"):
        run = read_synergy_kinetic(make_headerless_channel_file(sheet_names=(spelling,)))
        assert run.channel_names == ["OD600"], f"{spelling!r} did not resolve"


def test_a_derived_sheet_keyed_by_dose_is_not_mistaken_for_a_kinetic_block(tmp_path):
    """'UPRE1' holds Time against '0 mM', '0.1 mM' ... -- doses, not wells."""
    from openpyxl import Workbook

    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("UPRE1")
    ws.append([])
    ws.append(["Time", "0 mM", "0.1 mM", "0.5 mM"])
    for i in range(4):
        ws.append([i * 10, 3100 + i, 3050 + i, 3000 + i])
    od = wb.create_sheet("OD600")
    od.append(["Time", "A1", "A2"])
    for i, row in enumerate([[0.11, 0.12], [0.15, 0.16], [0.22, 0.24]]):
        od.append([f"0{i}:08:18", *row])
    path = tmp_path / "mixed.xlsx"
    wb.save(path)

    run = read_synergy_kinetic(path)

    assert run.channel_names == ["OD600"]


def test_time_is_still_parsed_when_the_channel_column_is_absent(make_headerless_channel_file):
    run = read_synergy_kinetic(make_headerless_channel_file())

    assert run.channel("OD600").times_h == pytest.approx(
        [8.3 / 60, 18.3 / 60, 68.3 / 60], abs=1e-2
    )


def test_the_temperature_trace_is_empty_rather_than_fabricated(make_headerless_channel_file):
    run = read_synergy_kinetic(make_headerless_channel_file())

    assert run.channel("OD600").temperature_c == []


def test_channels_from_separate_sheets_still_align(make_headerless_channel_file):
    run = read_synergy_kinetic(make_headerless_channel_file())

    aligned = run.aligned(reference="OD600")
    assert set(aligned.columns.get_level_values("channel")) == {"OD600", "mCitrine"}
    assert not aligned.isna().to_numpy().any()


class TestTimeIsNotAChannel:
    def test_a_column_literally_named_time_is_not_parsed_as_a_channel(self):
        from ystwin.plate.synergy import _parse_channel_header

        assert _parse_channel_header("Time") is None

    def test_the_degree_prefix_is_only_stripped_when_a_degree_sign_is_present(self):
        from ystwin.plate.synergy import _parse_channel_header

        assert _parse_channel_header("T° OD600:600") == ("OD600", "600")
        assert _parse_channel_header("Temperature") == ("Temperature", "")

    def test_a_second_time_column_does_not_create_a_phantom_channel(self, tmp_path):
        from openpyxl import Workbook

        from ystwin.plate.synergy import read_synergy_kinetic

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("OD600")
        ws.append(["Time", "Time", "A1", "A2"])
        for i, row in enumerate([[0.11, 0.12], [0.15, 0.16], [0.22, 0.24]]):
            ws.append([f"0{i}:08:18", f"0{i}:08:18", *row])
        path = tmp_path / "double_time.xlsx"
        wb.save(path)

        assert read_synergy_kinetic(path).channel_names == ["OD600"]


class TestDerivedSheets:
    """A 'for plotting' sheet interleaves summary columns among the wells."""

    def _workbook(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        raw = wb.create_sheet("OD600")
        raw.append(["Time", "A1", "A2", "A3"])
        derived = wb.create_sheet("OD600 (For Plotting)")
        derived.append(["Time", "A1", "A2", "Average", "A3"])
        for i, row in enumerate([[0.11, 0.12, 0.13], [0.15, 0.16, 0.17], [0.22, 0.24, 0.26]]):
            raw.append([f"0{i}:08:18", *row])
            derived.append([f"0{i}:08:18", row[0], row[1], sum(row) / 3, row[2]])
        path = tmp_path / "with_derived.xlsx"
        wb.save(path)
        return path

    def test_a_summary_column_among_the_wells_marks_a_sheet_as_derived(self, tmp_path):
        from ystwin.plate.synergy import read_synergy_kinetic

        run = read_synergy_kinetic(self._workbook(tmp_path), prefer_raw=False)
        by_sheet = {b.sheet: b.derived for b in run.blocks}

        assert by_sheet["OD600"] is False
        assert by_sheet["OD600 (For Plotting)"] is True

    def test_the_raw_sheet_wins_by_default_so_the_channel_stays_unambiguous(self, tmp_path):
        from ystwin.plate.synergy import read_synergy_kinetic

        run = read_synergy_kinetic(self._workbook(tmp_path))

        assert run.channel_names == ["OD600"]
        assert run.channel("OD600").derived is False

    def test_derived_blocks_are_still_reachable_when_asked_for(self, tmp_path):
        from ystwin.plate.synergy import read_synergy_kinetic

        run = read_synergy_kinetic(self._workbook(tmp_path), prefer_raw=False)

        assert len(run.blocks) == 2

    def test_a_derived_sheet_alone_is_kept_rather_than_discarded(self, tmp_path):
        from openpyxl import Workbook

        from ystwin.plate.synergy import read_synergy_kinetic

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("OD600 (For Plotting)")
        ws.append(["Time", "A1", "Average", "A2"])
        for i in range(3):
            ws.append([f"0{i}:08:18", 0.11 + i, 0.12 + i, 0.13 + i])
        path = tmp_path / "only_derived.xlsx"
        wb.save(path)

        assert read_synergy_kinetic(path).channel_names == ["OD600"]


class TestBlankSubtractedBlocks:
    """Processing is declared independently of the sign of the exported readings."""

    def _stacked(self, tmp_path):
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("OD600 (Raw)")
        ws.append(["Time", "A1", "A2"])
        for i, row in enumerate([[0.11, 0.12], [0.15, 0.16], [0.22, 0.24]]):
            ws.append([f"0{i}:08:18", *row])
        ws.append([])
        ws.append(["Time", "A1", "A2"])
        for i, row in enumerate([[0.02, 0.03], [-0.003, 0.07], [0.13, 0.15]]):
            ws.append([f"0{i}:08:18", *row])
        path = tmp_path / "stacked.xlsx"
        wb.save(path)
        return path

    def test_sign_does_not_declare_an_unknown_workbooks_processing_state(self, tmp_path):
        run = read_synergy_kinetic(self._stacked(tmp_path))

        assert [b.blank_subtracted for b in run.blocks] == [None, None]
        with pytest.raises(ValueError, match="correction-state"):
            run.raw_channel("OD600")

    def test_declaring_processing_does_not_change_channel_identity_or_values(self, tmp_path):
        import numpy as np

        path = self._stacked(tmp_path)
        undeclared = read_synergy_kinetic(path)
        run = read_synergy_kinetic(
            path, processing_states={"OD600[1]": False, "OD600[2]": True})

        assert run.channel_names == undeclared.channel_names
        assert [b.blank_subtracted for b in run.blocks] == [False, True]
        for before, after in zip(undeclared.blocks, run.blocks):
            assert np.array_equal(before.data.to_numpy(), after.data.to_numpy())
        assert run.raw_channel("OD600").channel == "OD600[1]"

    def test_a_declared_negative_raw_block_is_selectable_beside_positive_corrected_data(self, tmp_path):
        run = read_synergy_kinetic(
            self._stacked(tmp_path), processing_states={"OD600[1]": True, "OD600[2]": False})

        assert run.channel("OD600[1]").data.to_numpy().min() > 0
        assert run.raw_channel("OD600").channel == "OD600[2]"
        assert run.raw_channel("OD600").data.to_numpy().min() == -0.003

    def test_asking_for_a_raw_channel_that_has_none_is_refused(self, tmp_path):
        from openpyxl import Workbook

        from ystwin.plate.synergy import read_synergy_kinetic

        wb = Workbook()
        wb.remove(wb.active)
        ws = wb.create_sheet("OD600")
        ws.append(["Time", "A1", "A2"])
        for i, row in enumerate([[0.02, 0.03], [-0.003, 0.07], [0.13, 0.15]]):
            ws.append([f"0{i}:08:18", *row])
        path = tmp_path / "blanked_only.xlsx"
        wb.save(path)

        with pytest.raises(KeyError, match="already blank-subtracted"):
            read_synergy_kinetic(path, processing_states={"OD600": True}).raw_channel("OD600")

    def test_a_single_unblanked_channel_is_returned_by_raw_channel(
        self, make_headerless_channel_file
    ):
        from ystwin.plate.synergy import read_synergy_kinetic

        run = read_synergy_kinetic(
            make_headerless_channel_file(), processing_states={"OD600": False})

        assert run.raw_channel("OD600").channel == "OD600"


class TestChoosingBetweenRawCandidates:
    """Only declarations distinguish raw and corrected blocks, not coverage or ordering."""

    def _two_candidates(self, tmp_path, subset_first=True):
        from openpyxl import Workbook

        wb = Workbook()
        wb.remove(wb.active)
        wide = ["A1", "A2", "A3", "B1", "B2", "B3"]
        narrow = ["A1", "A2"]
        sheets = [("Blank mCitrine", narrow), ("mCitrine", wide)]
        if not subset_first:
            sheets.reverse()
        for name, wells in sheets:
            ws = wb.create_sheet(name)
            ws.append(["Time", *wells])
            for i in range(3):
                ws.append([f"0{i}:08:18", *[200 + 10 * i + 5 * j for j in range(len(wells))]])
        path = tmp_path / "candidates.xlsx"
        wb.save(path)
        return path

    def test_more_wells_do_not_resolve_ambiguous_correction_state(self, tmp_path):
        from ystwin.plate.synergy import read_synergy_kinetic

        run = read_synergy_kinetic(
            self._two_candidates(tmp_path), prefer_raw=False,
            processing_states={"mCitrine[1]": False, "mCitrine[2]": False})

        with pytest.raises(ValueError, match="ambiguous"):
            run.raw_channel("mCitrine")

    def test_sheet_order_does_not_decide_it(self, tmp_path):
        from ystwin.plate.synergy import read_synergy_kinetic

        first = read_synergy_kinetic(self._two_candidates(tmp_path, subset_first=True),
                                     prefer_raw=False)
        second = read_synergy_kinetic(self._two_candidates(tmp_path, subset_first=False),
                                      prefer_raw=False)

        for run in (first, second):
            with pytest.raises(ValueError, match="ambiguous"):
                run.raw_channel("mCitrine")

    def test_a_blank_subtracted_block_is_still_excluded_even_if_it_is_wider(self, tmp_path):
        from openpyxl import Workbook

        from ystwin.plate.synergy import read_synergy_kinetic

        wb = Workbook()
        wb.remove(wb.active)
        wide = wb.create_sheet("OD600")
        wide.append(["Time", "A1", "A2", "A3", "B1"])
        for i in range(3):
            wide.append([f"0{i}:08:18", 0.01, 0.2, 0.3, 0.4])
        narrow = wb.create_sheet("Raw OD")
        narrow.append(["Time", "A1", "A2"])
        for i in range(3):
            narrow.append([f"0{i}:08:18", 0.11 + i, 0.12 + i])
        path = tmp_path / "wide_blanked.xlsx"
        wb.save(path)

        run = read_synergy_kinetic(
            path, prefer_raw=False, processing_states={"OD600[1]": True, "OD600[2]": False})

        assert run.raw_channel("OD600").sheet == "Raw OD"


@pytest.mark.parametrize("reverse", [False, True])
def test_explicit_positive_corrected_block_cannot_be_selected_as_raw(reverse):
    from dataclasses import replace
    from pathlib import Path
    import pandas as pd
    from ystwin.plate.synergy import KineticBlock, SynergyRun

    data = pd.DataFrame({"A1": [100.0, 120.0, 140.0]}, index=[0.0, 1.0, 2.0])
    raw = KineticBlock("opaque-a", "mCitrine", "480,530", "not-a-state", False, False, data, [])
    corrected = replace(raw, channel="opaque-b", blank_subtracted=True, data=data + 10.0)
    blocks = (raw, corrected) if reverse else (corrected, raw)
    run = SynergyRun(Path("opaque.xlsx"), blocks)

    assert run.raw_channel("mCitrine").channel == "opaque-a"


@pytest.mark.parametrize("negative_in_continuation", [False, True])
def test_negative_values_do_not_override_a_raw_declaration_in_any_well_group(negative_in_continuation):
    from dataclasses import replace
    from pathlib import Path

    import pandas as pd

    from ystwin.plate.synergy import SynergyRun, _extract_blocks

    values = pd.DataFrame([
        ["Time", "A1"], ["0:00:00", 100.0], ["1:00:00", 120.0], ["2:00:00", 130.0],
        [None, None], [None, "H1"], [None, 5.0], [None, 2.0], [None, 1.0],
    ])
    values.iat[6 if negative_in_continuation else 1, 1] = -5.0
    block = _extract_blocks(values, "mCitrine")[0]

    assert block.blank_subtracted is None
    with pytest.raises(ValueError, match="correction-state"):
        SynergyRun(Path("opaque.xlsx"), (block,)).raw_channel("mCitrine")
    declared = replace(block, blank_subtracted=False)
    run = SynergyRun(Path("opaque.xlsx"), (declared,))
    assert run.raw_channel("mCitrine").data.to_numpy().min() == -5.0
    assert run.raw_channel("mCitrine").blank_subtracted is False


@pytest.mark.parametrize("declaration", [None, "False", 0])
def test_missing_or_nonboolean_correction_declaration_is_not_raw(declaration):
    from pathlib import Path
    import pandas as pd
    from ystwin.plate.synergy import KineticBlock, SynergyRun

    block = KineticBlock("opaque", "mCitrine", "480,530", "anything", False,
                         declaration, pd.DataFrame({"A1": [100.0]}), [])
    with pytest.raises(ValueError, match="correction-state"):
        SynergyRun(Path("opaque.xlsx"), (block,)).raw_channel("mCitrine")


@pytest.mark.parametrize("states", [True, {"OD600": "False"}, {"OD600": 0},
                                    {"OD600": None}, {"absent": False}])
def test_explicit_processing_declarations_require_known_channels_and_boolean_states(
        make_headerless_channel_file, states):
    with pytest.raises(ValueError, match="processing_states|correction-state"):
        read_synergy_kinetic(make_headerless_channel_file(), processing_states=states)


def test_a_partial_declaration_cannot_hide_an_unknown_candidate(tmp_path):
    path = TestBlankSubtractedBlocks()._stacked(tmp_path)
    run = read_synergy_kinetic(path, processing_states={"OD600[1]": False})

    with pytest.raises(ValueError, match="correction-state"):
        run.raw_channel("OD600")


@pytest.mark.parametrize("sheets", [("Raw OD", "BLANK OD"), ("BLANK OD", "Raw OD")])
def test_known_source_hash_and_exact_locators_supply_processing_not_name_or_order(
        make_headerless_channel_file, monkeypatch, sheets):
    import hashlib

    import pandas as pd

    from ystwin.plate import replay

    path = make_headerless_channel_file(sheet_names=sheets)
    digest = hashlib.sha256(path.read_bytes()).hexdigest()
    manifest = pd.DataFrame([
        {"source_sha256": digest, "sheet": sheet, "fluorophore": "OD600", "optics": "",
         "blank_subtracted": state, "export": "a_different_export_name.xlsx"}
        for sheet, state in (("Raw OD", "True"), ("BLANK OD", "False"))
    ])
    monkeypatch.setattr(replay, "load_manifest", lambda: manifest)

    run = read_synergy_kinetic(path)

    assert {b.sheet: b.blank_subtracted for b in run.blocks} == {"Raw OD": True, "BLANK OD": False}
    assert all(b.data.to_numpy().min() > 0 for b in run.blocks)
    assert run.raw_channel("OD600").sheet == "BLANK OD"
    renamed = path.rename(path.with_name("renamed.xlsx"))
    assert read_synergy_kinetic(renamed).raw_channel("OD600").sheet == "BLANK OD"
    with pytest.raises(ValueError, match="conflicts with source metadata"):
        read_synergy_kinetic(renamed, processing_states={run.raw_channel("OD600").channel: True})


def test_a_familiar_filename_cannot_classify_different_workbook_bytes(make_headerless_channel_file):
    path = make_headerless_channel_file(name="20260803_ER&oxidativestress_Replicate3.xlsx")

    run = read_synergy_kinetic(path)

    assert all(b.blank_subtracted is None for b in run.blocks)
    with pytest.raises(ValueError, match="correction-state"):
        run.raw_channel("OD600")


def test_a_source_locator_matching_two_blocks_does_not_assign_one_declaration_twice(tmp_path, monkeypatch):
    from ystwin.plate import replay

    monkeypatch.setattr(replay, "_processing_states_for_source",
                        lambda digest: {("OD600 (Raw)", "OD600", ""): False})
    path = TestBlankSubtractedBlocks()._stacked(tmp_path)
    run = read_synergy_kinetic(path)

    assert [b.blank_subtracted for b in run.blocks] == [None, None]
    with pytest.raises(ValueError, match="correction-state"):
        run.raw_channel("OD600")
    explicit = read_synergy_kinetic(path, processing_states={"OD600[1]": False, "OD600[2]": True})
    assert explicit.raw_channel("OD600").channel == "OD600[1]"
