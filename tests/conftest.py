"""Shared fixtures. Synthetic plate files mimic real Synergy H1 exports.

Private original workbooks are optional and locate data through :mod:`ystwin.paths`, not
by counting directories upwards. Their committed public measurements remain mandatory,
as do public models and their source-check runtimes in strict CI. A clone need not have
seen a plate reader, but missing public coverage must never become a green skip.
"""
from __future__ import annotations

import json
import os
import pathlib
import subprocess
import sys
from copy import deepcopy

import numpy as np
import pytest

# Test the checkout, not whatever `pip install -e` last resolved. Without this a
# bisect across commits silently measures the working tree at every step.
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parents[1] / "src"))
from openpyxl import Workbook

from ystwin import paths

_REPO_ROOT = pathlib.Path(__file__).resolve().parents[1]
_ENV_FILE = _REPO_ROOT / ".ystwin.env"


def _in_ci():
    return any(os.environ.get(name, "").lower() in {"1", "true"}
               for name in ("CI", "GITHUB_ACTIONS"))


def _requires_public_inputs(config):
    return (config.getoption("--require-public-inputs", default=False)
            or (_in_ci() and config.rootpath.resolve() == _REPO_ROOT))


def pytest_addoption(parser):
    parser.addoption("--require-public-inputs", action="store_true",
                     help="fail on missing mandatory public inputs, source runtimes, or unclassified skips; private originals remain optional")


def _private_input_for(item, catalogue):
    try:
        filename = item.path.resolve().relative_to(_REPO_ROOT).as_posix()
    except ValueError:
        return None
    ancestry = {node.nodeid for node in item.listchain()}
    for record in catalogue:
        fixtures = set(record["fixtures"]) | set(record["fixture_selectors"].get(filename, []))
        if fixtures.intersection(item.fixturenames) or ancestry.intersection(record["test_classes"]):
            return record
    return None


def pytest_collection_modifyitems(items):
    manifest = json.loads((_REPO_ROOT / "data/public_inputs.json").read_text(encoding="utf-8"))
    for item in items:
        record = _private_input_for(item, manifest["optional_private_inputs"])
        if record is not None:
            item.add_marker(pytest.mark.private_data(record["environment"], reason=record["reason"]))
            item.user_properties.append(("input_scope", "optional_private"))
            item.user_properties.append(("private_input", record["environment"]))


def _verify_public_inputs():
    return subprocess.run(
        [sys.executable, str(_REPO_ROOT / "scripts/provision_public_inputs.py"), "--verify-only"],
        cwd=_REPO_ROOT, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, check=False)


def _verify_public_source_runtimes():
    return subprocess.run(
        [sys.executable, str(_REPO_ROOT / "scripts/verify_quality.py"), "--check-source-runtimes"],
        cwd=_REPO_ROOT, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True, check=False)


def pytest_collection_finish(session):
    if _requires_public_inputs(session.config) and not session.config.option.collectonly:
        for label, verify in (("public input", _verify_public_inputs),
                              ("public source runtime", _verify_public_source_runtimes)):
            result = verify()
            if result.returncode != 0:
                pytest.exit(f"Mandatory {label} verification failed (exit {result.returncode}):\n{result.stdout}",
                            returncode=1)


@pytest.hookimpl(hookwrapper=True)
def pytest_make_collect_report(collector):
    outcome = yield
    report = outcome.get_result()
    if report.skipped and _requires_public_inputs(collector.config):
        report.outcome = "failed"
        report.longrepr = f"A mandatory test module skipped during CI collection: {report.longrepr}"


@pytest.hookimpl(hookwrapper=True)
def pytest_runtest_makereport(item, call):
    outcome = yield
    report = outcome.get_result()
    if not report.skipped:
        return
    private = item.get_closest_marker("private_data")
    if private is not None:
        filename, line, reason = report.longrepr
        report.longrepr = (filename, line, f"optional private input {private.args[0]}: {reason}")
    elif _requires_public_inputs(item.config):
        report.outcome = "failed"
        report.longrepr = f"{item.nodeid}: mandatory CI coverage cannot skip: {report.longrepr}"


def _load_local_env() -> None:
    """Read ``.ystwin.env`` if the machine has one, without overriding a real variable.

    **Why this exists.** `ystwin.paths` deliberately has no machine-local defaults: a path
    found by convention makes a result reproducible on one machine only, and both audit
    gates now refuse one. The cost was that a machine which HAS the wet-lab data still ran
    the suite with the data invisible, because eleven separate variables had to be exported
    by hand and nothing said so. Fourteen tests skipped on the machine holding every file
    they needed.

    So the paths live in an untracked file next to the checkout, `.ystwin.env`, listed in
    `.gitignore` and templated by `.ystwin.env.example`. Nothing tracked gains a path,
    nothing resolves by convention, and a machine with the data can see it. An existing
    environment variable always wins, so CI and a one-off `YSTWIN_... = ... pytest` are
    unaffected.
    """
    if _in_ci() or not _ENV_FILE.is_file():
        return
    for line in _ENV_FILE.read_text(encoding="utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        name, _, value = line.partition("=")
        name, value = name.strip(), value.strip().strip('"').strip("'")
        if not name.startswith("YSTWIN_"):
            continue
        os.environ.setdefault(name, value)


_load_local_env()


def _public_unavailable(message):
    if _in_ci():
        pytest.fail(f"mandatory public input: {message}", pytrace=False)
    pytest.skip(message)


def _real_export(name: str) -> pathlib.Path:
    """One named Synergy **workbook**, or a skip naming the variable that would find it.

    The workbook, not its numbers. Everything reached through this fixture tests the parser
    against a real ``.xlsx`` -- sheet layout, block boundaries, well ordering, the elapsed
    time column -- and the committed text under ``data/plates`` cannot stand in for that,
    because it is what the parser produced. Replaying it here would turn a test of the
    reader into a test of the replay.

    So these skip when the workbooks are absent, which is now the ordinary case:
    ``paths.igem_results()`` no longer falls back to a sibling directory, because the
    numbers those workbooks carry are committed and the convention made every result
    computed through it reproducible on one machine only. What is lost by skipping is
    parser coverage against a real file; what is not lost is any number, because
    ``plate_readings.py --export`` verifies the committed text against these same workbooks
    value-for-value and refuses to leave an export in place that does not match.
    """
    results = paths.igem_results()
    if results is None:
        pytest.skip(
            "the 2026-07 Synergy workbooks are not here. Their measurements are committed "
            "under data/plates and every result is reproducible without them; set "
            "YSTWIN_IGEM_RESULTS to also exercise the parser against the .xlsx itself")
    path = results / name
    if not path.exists():
        pytest.skip(f"real wet-lab file not present: {path}")
    return path


def _write_kinetic_sheet(ws, channel_header, wells, times, values, temps):
    """Lay out one Synergy-style kinetic block starting at the sheet's current row."""
    if channel_header is None:
        # Some exports drop the temperature/channel column entirely; the channel is
        # then only identifiable from the sheet name.
        ws.append(["Time", *wells])
        for t, row in zip(times, values):
            ws.append([t, *row])
        return
    ws.append(["Time", channel_header, *wells])
    for t, temp, row in zip(times, temps, values):
        ws.append([t, temp, *row])


@pytest.fixture
def make_synergy_file(tmp_path):
    """Build a minimal two-channel Synergy export with an optional preamble offset."""

    def _make(name="plate.xlsx", lead_blank_rows=0, wells=("A1", "A2", "B1"), pad_rows=0,
              od_times=("00:07:25", "00:17:25", "01:07:25"),
              second_reporter_read=False):
        wb = Workbook()
        od = wb.active
        od.title = "OD600"
        for _ in range(lead_blank_rows):
            od.append([])
        _write_kinetic_sheet(
            od,
            "T° OD600:600",
            wells,
            list(od_times),
            [[0.10, 0.20, 0.30], [0.12, 0.24, 0.33], [0.20, 0.40, 0.50]],
            [30.0, 30.1, 30.0],
        )
        fl = wb.create_sheet("Fluorescence")
        for _ in range(lead_blank_rows):
            fl.append([])
        _write_kinetic_sheet(
            fl,
            "T° mCitrine:480,530",
            wells,
            ["00:08:15", "00:18:15", "01:08:15"],
            [[1000, 2000, 3000], [1200, 2400, 3300], [2000, 4000, 5000]],
            [30.0, 30.0, 30.1],
        )
        # Real exports pad the block out to the protocol's full length with
        # "00:00:00" timestamps and empty wells.
        for sheet in (od, fl):
            for _ in range(pad_rows):
                sheet.append(["00:00:00", None, *[None] * len(wells)])
        if second_reporter_read:
            fl2 = wb.create_sheet("Fluorescence2")
            _write_kinetic_sheet(
                fl2,
                "T° mCitrine:480,530[2]",
                wells,
                ["00:09:05", "00:19:05", "01:09:05"],
                [[80, 160, 240], [96, 192, 264], [160, 320, 400]],
                [30.0, 30.0, 30.0],
            )
            for _ in range(pad_rows):
                fl2.append(["00:00:00", None, *[None] * len(wells)])
        path = tmp_path / name
        wb.save(path)
        return path

    return _make


@pytest.fixture(scope="session")
def real_er_prelim():
    return _real_export("20260701_ER_preliminary (RAW).xlsx")


@pytest.fixture(scope="session")
def real_er_stress():
    return _real_export("20260709_ER_stress_1st (RAW).xlsx")


@pytest.fixture(scope="session")
def thermo_data():
    """The eQuilibrator formation energies, or a skip naming the variable that finds them.

    Three test files built this by calling ``ThermodynamicData.load()`` in a local
    fixture, which raises when the asset is absent. That put 54 tests into ERROR on a
    checkout without it, against this module's stated contract -- and an error reads as a
    broken test rather than an absent input.
    """
    from ystwin.bridge.thermodynamic import ThermodynamicData

    try:
        return ThermodynamicData.load()
    except (FileNotFoundError, SystemExit) as exc:
        _public_unavailable(f"thermodynamic data not present; set YSTWIN_THERMO ({exc})")


def _atp_export(name: str) -> pathlib.Path:
    """One named ATP-sensor export, or a skip naming the variable that would find it."""
    directory = paths.atp_sensor_plates()
    if directory is None:
        pytest.skip("ATP-sensor exports not present; set YSTWIN_ATP_SENSOR")
    path = directory / name
    if not path.exists():
        pytest.skip(f"real wet-lab file not present: {path}")
    return path


@pytest.fixture(scope="session")
def atp_sensor_icl():
    """The 2026-08-05 plate: the ICL-UAS build, one sensor over a carbon grid."""
    return _atp_export("20260805_ICL.xlsx")


@pytest.fixture(scope="session")
def atp_sensor_acs_icl():
    """The 2026-08-09 plate: two builds, rows A-C and D-F, same carbon grid."""
    return _atp_export("20260809_ACS_ICL.xlsx")


def _copy_model(template):
    import swiglpk as glpk
    from optlang import glpk_interface

    solver = template.solver
    solver.update()
    if isinstance(solver, glpk_interface.Model):
        problem = glpk.glp_create_prob()
        glpk.glp_copy_prob(problem, solver.problem, glpk.GLP_ON)
        copied_solver = type(solver)(problem=problem)
        for index, variable in enumerate(copied_solver.variables, 1):
            kind = glpk.glp_get_col_type(solver.problem, index)
            lower = (glpk.glp_get_col_lb(solver.problem, index)
                     if kind in (glpk.GLP_LO, glpk.GLP_DB, glpk.GLP_FX) else None)
            upper = (glpk.glp_get_col_ub(solver.problem, index)
                     if kind in (glpk.GLP_UP, glpk.GLP_DB, glpk.GLP_FX) else None)
            if (variable.lb, variable.ub) != (lower, upper):
                variable.set_bounds(lower, upper)
        copied_solver.objective = solver.interface.Objective.clone(
            solver.objective, model=copied_solver)
        glpk.glp_std_basis(copied_solver.problem)
        model = deepcopy(template, {id(solver): copied_solver})
    else:
        model = deepcopy(template)
        for variable in template.variables:
            model.variables[variable.name].set_bounds(variable.lb, variable.ub)
    for group in model.groups:
        group._model = model
    configuration = solver.configuration
    model.solver.configuration = type(configuration).clone(configuration, problem=model.solver)
    return model


@pytest.fixture(scope="session")
def model_copy():
    return _copy_model


@pytest.fixture(scope="session")
def _yeast_gem_template():
    """Yeast9 GSMM. Session-scoped: SBML parsing costs about a second."""
    path = paths.yeast_gem()
    if path is None:
        _public_unavailable("yeast-GEM not present; set YSTWIN_YEAST_GEM to yeast-GEM.xml")
    import cobra

    return cobra.io.read_sbml_model(str(path))


@pytest.fixture(scope="session")
def yeast_gem_factory(_yeast_gem_template):
    def make():
        return _copy_model(_yeast_gem_template)
    return make


@pytest.fixture
def yeast_gem(yeast_gem_factory):
    return yeast_gem_factory()


@pytest.fixture
def carotenoid_model(yeast_gem):
    from ystwin.fba.carotenoid import add_beta_carotene_pathway

    return add_beta_carotene_pathway(yeast_gem)


@pytest.fixture(scope="session")
def _ec_yeast_gem_template():
    """GECKO enzyme-constrained yeast model (batch). Note: built on yeast-GEM 8.3.4."""
    path = paths.ec_yeast_gem()
    if path is None:
        _public_unavailable("ecYeastGEM not present; set YSTWIN_EC_YEAST_GEM to ecYeastGEM_batch.xml")
    import cobra

    return cobra.io.read_sbml_model(str(path))


@pytest.fixture(scope="session")
def ec_yeast_gem_factory(_ec_yeast_gem_template):
    def make():
        return _copy_model(_ec_yeast_gem_template)
    return make


@pytest.fixture
def ec_yeast_gem(ec_yeast_gem_factory):
    return ec_yeast_gem_factory()


@pytest.fixture
def make_headerless_channel_file(tmp_path):
    """Export whose sheets carry no temperature/channel column, as in the 2026-08 plates."""

    def _make(name="plate.xlsx", sheet_names=("OD600", "mCitrine"), lead_blank_rows=2,
              wells=("A1", "A2", "B1")):
        wb = Workbook()
        wb.remove(wb.active)
        payload = {
            "OD600": [[0.11, 0.12, 0.13], [0.14, 0.16, 0.17], [0.22, 0.25, 0.27]],
            "mCitrine": [[350, 360, 340], [420, 435, 410], [640, 660, 630]],
        }
        for sheet in sheet_names:
            ws = wb.create_sheet(sheet)
            for _ in range(lead_blank_rows):
                ws.append([])
            _write_kinetic_sheet(
                ws, None, wells,
                ["00:08:18", "00:18:18", "01:08:18"],
                payload.get(sheet, payload["OD600"]), None,
            )
        path = tmp_path / name
        wb.save(path)
        return path

    return _make


@pytest.fixture(scope="session")
def plate_text():
    """Elapsed times from one committed export, as hours.

    Reads `data/plates`, which is tracked, so this never skips -- the geometry constants in
    `analysis/estimator_accuracy.py` are checked against the plates themselves rather than
    against a remembered number.
    """
    import pandas as pd

    from ystwin.plate import replay

    frame = pd.read_csv(
        replay.committed_dir()
        / "20260722_ERandOxidativeStress_NewProtocol_ANALYSED__mCitrine_1.csv")

    def hours(text: str) -> float:
        h, m, s = (float(part) for part in text.split(":"))
        return h + m / 60.0 + s / 3600.0

    return np.array([hours(value) for value in frame.elapsed_hms])
